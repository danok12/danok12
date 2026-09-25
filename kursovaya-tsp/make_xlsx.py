"""Книга Excel: исходные отметки, точки нулевых работ, объёмы ПВ/ПН (с формулами)."""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.comments import Comment
import raschet as R

FONT = "GOST type B"            # ГОСТ 2.304, тип Б (файл GOST_B.TTF)
thin = Side(style="thin")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
F_IN = Font(name=FONT, size=12, color="0000FF")          # исходные данные
F_LINK = Font(name=FONT, size=12, color="008000")        # ссылка на другой лист
F_CALC = Font(name=FONT, size=12)
F_B = Font(name=FONT, size=12, bold=True)
F_T = Font(name=FONT, size=14, bold=True)
YEL = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="DDEBF7")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

wb = Workbook()
S1 = "Исходные данные"
ws = wb.active
ws.title = S1


def cell(sh, ref, val, font=F_CALC, fmt=None, al=CEN, box=True, fill=None):
    c = sh[ref]
    c.value = val
    c.font = font
    c.alignment = al
    if fmt: c.number_format = fmt
    if box: c.border = BOX
    if fill: c.fill = fill
    return c


# ---------- лист 1: исходные данные ----------
ws["A1"] = "Курсовая работа «Технологическая карта на производство земляных работ». Вариант 2"
ws["A1"].font = F_T
cell(ws, "A3", "Сторона квадрата сетки a, м", F_CALC, al=LEFT)
cell(ws, "B3", R.A, F_IN)
cell(ws, "A4", "Коэффициент остаточного разрыхления Kор", F_CALC, al=LEFT)
cell(ws, "B4", 1.04, F_IN, "0.00", fill=YEL)
ws["B4"].comment = Comment(
    "Супесь: остаточное разрыхление 3–5 % (табл. П2.1 методички), Kор = 1,03…1,05. "
    "Принято 1,04 — по указанию преподавателя брать в этом диапазоне.", "расчёт")
cell(ws, "A5", "Грунт", F_CALC, al=LEFT)
cell(ws, "B5", "супесь", F_IN)
cell(ws, "A6", "Число квадратов", F_CALC, al=LEFT)
cell(ws, "B6", R.NX * R.NY, F_IN)
ws["D3"] = "Редактировать можно только синие ячейки (жёлтая — принятое допущение)."
ws["D4"] = "Рабочая отметка h = Hкр − Hчёрн: «−» — выемка (ПВ), «+» — насыпь (ПН)."
ws["D5"] = "Отметки — приложение 1 методички, вариант 2. Вершины нумеруются построчно слева направо, сверху вниз."
for r in ("D3", "D4", "D5"):
    ws[r].font = Font(name=FONT, size=11, italic=True)

cell(ws, "A8", "№ вершины", F_B, fill=HEAD)
cell(ws, "B8", "h, м", F_B, fill=HEAD)
HREF = {}
for k, v in enumerate(R.H, start=1):
    r = 8 + k
    cell(ws, f"A{r}", f"h{k}")
    cell(ws, f"B{r}", float(v), F_IN, "0.00")
    HREF[k] = f"'{S1}'!$B${r}"

ws["D8"] = "Схема сетки (вид в плане, квадрат 100×100 м)"
ws["D8"].font = F_B
cols = "DEFGHI"
for j in range(R.NY + 1):
    for i in range(R.NX + 1):
        n = R.vnum(i, j)
        cell(ws, f"{cols[i]}{9 + 2 * j}", f"h{n}", Font(name=FONT, size=9, color="808080"))
        cell(ws, f"{cols[i]}{10 + 2 * j}", f"=B{8 + n}", F_CALC, "0.00")
ws.column_dimensions["A"].width = 44
ws.column_dimensions["B"].width = 10
for c in cols: ws.column_dimensions[c].width = 9

# ---------- лист 2: ЛНР ----------
S2 = "ЛНР"
w2 = wb.create_sheet(S2)
w2["A1"] = "Точки нулевых работ на сторонах квадратов"
w2["A1"].font = F_T
w2["A2"] = "x1 = a·|h1| / (|h1| + |h2|) — расстояние от вершины 1 до точки нуля;  x2 = a − x1"
w2["A2"].font = Font(name=FONT, size=12, italic=True)
hdr = ["№ т.0", "Сторона", "Вершина 1", "h1, м", "Вершина 2", "h2, м",
       "x1 (от верш. 1), м", "x2 (от верш. 2), м"]
for k, t in enumerate(hdr):
    cell(w2, f"{'ABCDEFGH'[k]}4", t, F_B, fill=HEAD)
ZROW = {}
zp = sorted(R.ZP, key=lambda z: (z["p1"], z["p2"]))
for k, z in enumerate(zp, start=1):
    r = 4 + k
    ZROW[(z["p1"], z["p2"])] = r
    horiz = z["p2"] == z["p1"] + 1
    cell(w2, f"A{r}", k)
    cell(w2, f"B{r}", "горизонтальная" if horiz else "вертикальная")
    cell(w2, f"C{r}", f"h{z['p1']}")
    cell(w2, f"D{r}", f"={HREF[z['p1']]}", F_LINK, "0.00")
    cell(w2, f"E{r}", f"h{z['p2']}")
    cell(w2, f"F{r}", f"={HREF[z['p2']]}", F_LINK, "0.00")
    cell(w2, f"G{r}", f"='{S1}'!$B$3*ABS(D{r})/(ABS(D{r})+ABS(F{r}))", F_CALC, "0.00")
    cell(w2, f"H{r}", f"='{S1}'!$B$3-G{r}", F_CALC, "0.00")
for c, wd in zip("ABCDEFGH", (7, 16, 11, 9, 11, 9, 18, 18)):
    w2.column_dimensions[c].width = wd

# ---------- лист 3: объёмы ----------
S3 = "Объёмы"
w3 = wb.create_sheet(S3)
w3["A1"] = "Объёмы планировочных работ (метод квадратов)"
w3["A1"].font = F_T
w3["A2"] = "Vi = Fi · hср,i;   hср,i = (|h1| + |h2| + … + |hn|) / n, в точках нуля h = 0"
w3["A2"].font = Font(name=FONT, size=12, italic=True)
A2 = f"'{S1}'!$B$3^2"


def dist(v, e):
    """Ссылка на расстояние от вершины v до точки нуля на стороне e."""
    r = ZROW[e]
    return f"'{S2}'!$G${r}" if v == e[0] else f"'{S2}'!$H${r}"


FCELL = {}      # имя фигуры -> адрес ячейки F


def describe(f):
    vs = [lb[1] for lb in f["labels"] if lb[0] == "v"]
    zs = [lb[1] for lb in f["labels"] if lb[0] == "z"]
    kind = {4: "квадрат", 1: "треугольник", 2: "трапеция", 3: "пятиугольник"}[len(vs)]
    txt = ", ".join(f"h{v}" for v in vs)
    if zs:
        txt += "; т.0 на " + ", ".join(f"h{a}–h{b}" for a, b in zs)
    return kind, vs, zs, txt


def area_formula(f, vs, zs, other_ref):
    if len(vs) == 4:
        return f"={A2}"
    if len(vs) == 1:
        v = vs[0]
        return f"={dist(v, zs[0])}*{dist(v, zs[1])}/2"
    if len(vs) == 2:
        d = [dist(next(v for v in vs if v in e), e) for e in zs]
        return f"=({d[0]}+{d[1]})/2*'{S1}'!$B$3"
    return f"={A2}-{other_ref}"


def hcp_formula(f):
    parts = [f"ABS({HREF[lb[1]]})" if lb[0] == "v" else "0" for lb in f["labels"]]
    return f"=({'+'.join(parts)})/{len(parts)}"


# Одна таблица, как на доске: верхняя строка шапки делит её на ПВ и ПН,
# строки - по квадратам 1..15, где части нет - прочерк.
R_GRP, R_HDR = 4, 5
R1 = R_HDR + 1
NSQ = R.NX * R.NY
RS = R1 + NSQ                       # строка итогов
COLS = {-1: "ABCD", 1: "EFGHI"}     # ПВ: № F hср V;  ПН: № F hср V Vгр
for sq in range(1, NSQ + 1):
    for g in R.FIGS:
        if g["sq"] == sq:
            FCELL[g["name"]] = f"${COLS[g['sign']][1]}${R1 + sq - 1}"

w3.merge_cells(f"A{R_GRP}:D{R_GRP}")
w3.merge_cells(f"E{R_GRP}:I{R_GRP}")
cell(w3, f"A{R_GRP}", "Планировочная выемка (ПВ)", F_B, fill=HEAD)
cell(w3, f"E{R_GRP}", "Планировочная насыпь (ПН)", F_B, fill=HEAD)
for c in "BCDFGHI":
    w3[f"{c}{R_GRP}"].border = BOX
hdr = {-1: ["№ фигуры", "Fi, м²", "hср, м", "Vi, м³"],
       1: ["№ фигуры", "Fi, м²", "hср, м", "Vi, м³", "Vгр = Vi / Kор, м³"]}
for sgn in (-1, 1):
    for c, t in zip(COLS[sgn], hdr[sgn]):
        cell(w3, f"{c}{R_HDR}", t, F_B, fill=HEAD)

for sq in range(1, NSQ + 1):
    r = R1 + sq - 1
    for sgn in (-1, 1):
        cn = COLS[sgn]
        part = [g for g in R.FIGS if g["sq"] == sq and g["sign"] == sgn]
        if not part:
            for c in cn:
                cell(w3, f"{c}{r}", "—")
            continue
        f = part[0]
        kind, vs, zs, txt = describe(f)
        sibling = [g for g in R.FIGS if g["sq"] == sq and g is not f]
        other = FCELL[sibling[0]["name"]] if sibling else None
        cell(w3, f"{cn[0]}{r}", f["name"], F_B)
        w3[f"{cn[0]}{r}"].comment = Comment(f"{kind}: {txt}", "расчёт")
        cell(w3, f"{cn[1]}{r}", area_formula(f, vs, zs, other), F_CALC, "0.00")
        cell(w3, f"{cn[2]}{r}", hcp_formula(f), F_CALC, "0.000")
        cell(w3, f"{cn[3]}{r}", f"={cn[1]}{r}*{cn[2]}{r}", F_CALC, "0.00")
        if sgn > 0:
            cell(w3, f"{cn[4]}{r}", f"={cn[3]}{r}/'{S1}'!$B$4", F_CALC, "0.00")

for sgn in (-1, 1):
    cn = COLS[sgn]
    cell(w3, f"{cn[0]}{RS}", "Σ", F_B)
    cell(w3, f"{cn[2]}{RS}", "")
    for c in (cn[1], cn[3]) + ((cn[4],) if sgn > 0 else ()):
        cell(w3, f"{c}{RS}", f"=SUM({c}{R1}:{c}{RS - 1})", F_B, "0.00")
TOT = {-1: ("B", "D"), 1: ("F", "H", "I")}

rb = RS + 3
w3[f"A{rb - 1}"] = "Сопоставление объёмов"
w3[f"A{rb - 1}"].font = F_B
rows = [
    ("Площадь ПВ + ПН, м² (проверка)", f"=B{RS}+F{RS}", "0.00"),
    ("Площадь площадки 5 × 3 × a², м²", f"='{S1}'!$B$6*{A2}", "0.00"),
    ("ΣVпв — объём выемки, м³", f"=D{RS}", "0.00"),
    ("ΣVпн — объём насыпи, м³", f"=H{RS}", "0.00"),
    ("ΣVгр = ΣVпн / Kор — грунта в плотном теле на насыпь, м³", f"=I{RS}", "0.00"),
    ("Разница ΣVгр − ΣVпв, м³ («+» — не хватает грунта)", f"=F{rb + 4}-F{rb + 2}", "0.00"),
    ("Разница в % от ΣVпв", f"=F{rb + 5}/F{rb + 2}", "0.0%"),
]
for k, (t, fml, fmt) in enumerate(rows):
    w3.merge_cells(f"A{rb + k}:E{rb + k}")
    cell(w3, f"A{rb + k}", t, F_CALC, al=LEFT)
    for c in "BCDE":
        w3[f"{c}{rb + k}"].border = BOX
    cell(w3, f"F{rb + k}", fml, F_B, fmt)
note = rb + len(rows) + 1
w3[f"A{note}"] = ("Разницу покрывают в сводном балансе грунта: грунт из котлована, "
                  "завоз из резерва или поправка отметок Δh — следующий этап.")
w3[f"A{note}"].font = Font(name=FONT, size=11, italic=True)
w3.merge_cells(f"A{note}:I{note}")
w3[f"A{note}"].alignment = LEFT
w3.row_dimensions[note].height = 32
for c, wd in zip("ABCDEFGHI", (10, 11, 9, 11, 10, 11, 9, 11, 13)):
    w3.column_dimensions[c].width = wd
w3.row_dimensions[R_HDR].height = 32

for sh in wb.worksheets:
    sh.sheet_view.zoomScale = 110
    sh.page_setup.paperSize = sh.PAPERSIZE_A4
    sh.page_setup.orientation = "landscape"
    sh.sheet_properties.pageSetUpPr.fitToPage = True
    sh.page_setup.fitToWidth = 1
    sh.page_setup.fitToHeight = 1          # каждый лист печатается на одну страницу

out = "Объёмы_планировки_вариант2.xlsx"
wb.save(out)
print(out, "итоги в строке", RS, "сравнение с", rb)
