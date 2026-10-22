# -*- coding: utf-8 -*-
"""Книга Excel «Таблицы расчёта — ВиВ вариант 3»: все таблицы пояснительной записки, связанные формулами.

Исходные данные (жёлтые) -> расходы -> гидравлика В1 (Шевелёв) -> ввод (Колбрук–Уайт) -> счётчики -> напор ->
насос; К1: стояки, горизонтальные участки, отметки дворовой сети. α — по табл. Б.2 линейной интерполяцией (как у
преподавателя). Табличные величины, для которых в Excel нужен итерационный подбор (K_s по табл. 5.1, h/d и V по
Павловскому — таблицы Лукиных, напор насосов по кривым каталога), внесены числами из calc/raschet_v3.py.

Запуск: python sozdat_tablicy_kursovogo.py  ->  «Таблицы расчёта — ВиВ вариант 3.xlsx» рядом со скриптом.
"""
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
KURS = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(KURS, "calc"))
from normy import A2_ZHILYE  # noqa: E402

REZ = json.load(open(os.path.join(KURS, "calc", "rezultaty_v3.json"), encoding="utf-8"))
NAS = json.load(open(os.path.join(KURS, "calc", "nasos_v3.json"), encoding="utf-8"))
GP = json.load(open(os.path.join(KURS, "genplan", "genplan_v3.json"), encoding="utf-8"))
B2 = sorted((float(k), v) for k, v in json.load(open(os.path.join(KURS, "calc", "sp30_pril_B.json"),
                                                     encoding="utf-8"))["b2"].items())
NORMA = A2_ZHILYE["ванны от 1500 мм с душами"]
IZ = REZ["iznach"]
OUT = os.path.join(HERE, "Таблицы расчёта — ВиВ вариант 3.xlsx")

TNR = "Times New Roman"
f_b = Font(name=TNR, size=11)
f_bold = Font(name=TNR, size=11, bold=True)
f_in = Font(name=TNR, size=11, color="0000FF")
f_title = Font(name=TNR, size=13, bold=True)
f_note = Font(name=TNR, size=9, italic=True, color="595959")
fill_in = PatternFill("solid", fgColor="FFF2A8")
fill_h = PatternFill("solid", fgColor="F2F2F2")
fill_res = PatternFill("solid", fgColor="E2EFDA")
thin = Side(style="thin", color="808080")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEV = Alignment(horizontal="left", vertical="center", wrap_text=True)

wb = Workbook()
NAMES = {}


def imya(nm, ws, ref):
    NAMES[nm] = "'%s'!%s" % (ws.title, ref.replace("", "") if False else "$" + re.sub(r"(\d+)", r"$\1", ref))


def yach(ws, ref, val, font=f_b, fill=None, fmt=None, al=None, border=True):
    c = ws[ref]
    c.value = val
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if al is None and font is f_note:
        al = Alignment(horizontal="left", vertical="top", wrap_text=False)   # примечание — в строку
    c.alignment = al or (LEV if isinstance(val, str) and not val.startswith("=") else CEN)
    if border:
        c.border = box
    return c


def list_(nazv, zag, shir):
    ws = wb.create_sheet(nazv)
    ws.sheet_view.zoomScale = 100
    yach(ws, "A1", zag, f_title, border=False, al=Alignment(horizontal="left", vertical="center"))
    for j, w in enumerate(shir):
        ws.column_dimensions[chr(ord("A") + j)].width = w
    return ws


def shapka(ws, row, cols):
    for j, h in enumerate(cols):
        yach(ws, "%s%d" % (chr(ord("A") + j), row), h, f_bold, fill_h, al=CEN)
    ws.row_dimensions[row].height = 32


def alfa(x):
    """α по табл. Б.2 линейной интерполяцией (NP < 0,015 — 0,2)."""
    m = "MATCH(%s,B2_NP,1)" % x
    return ("IF(%s<0.015,0.2,INDEX(B2_A,%s)+(%s-INDEX(B2_NP,%s))*(INDEX(B2_A,%s+1)-INDEX(B2_A,%s))"
            "/(INDEX(B2_NP,%s+1)-INDEX(B2_NP,%s)))" % (x, m, x, m, m, m, m, m))


# ---------------- 0. содержание ----------------
w0 = wb.active
w0.title = "Содержание"
w0.column_dimensions["A"].width = 30
w0.column_dimensions["B"].width = 95
yach(w0, "A1", "Таблицы расчёта курсового проекта «Водоснабжение и водоотведение жилого здания», вариант 3",
     f_title, border=False, al=Alignment(horizontal="left"))
yach(w0, "A2", "Буев Даниил Романович, ИГЭС 3-8. Номера таблиц — как в пояснительной записке.", f_note, border=False,
     al=Alignment(horizontal="left"))
SODERZH = [
    ("1.1 Исходные", "Таблица 1.1. Исходные данные (задание) и принятые допущения — жёлтые ячейки можно менять"),
    ("1.2 Нормы", "Таблица 1.2. Нормы расхода воды (табл. А.2 СП 30.13330.2020)"),
    ("1.3 Расходы", "Таблица 1.3. Расчётные расходы воды: U, N, суточные, среднечасовые, секундные, часовые"),
    ("1.4 Гидравлика В1", "Таблица 1.4. Гидравлический расчёт сети В1 по диктующему направлению (Шевелёв)"),
    ("1.5 Ввод", "Расчёт ввода (ПЭ100 SDR17, Колбрук–Уайт)"),
    ("1.6 Счётчики", "Таблица 1.5. Подбор счётчиков воды"),
    ("1.7 Напор", "Требуемый напор, напор установки, регуляторы давления по этажам"),
    ("1.8 Насос", "Таблица 1.6. Подбор насоса Wilo MVL (напор по кривым каталога)"),
    ("2.1 Стояки К1", "Таблица 2.1. Пропускная способность стояков К1"),
    ("2.2 Горизонт. К1", "Таблица 2.2. Расчёт горизонтальных участков К1 и дворовой сети"),
    ("2.3 Отметки", "Таблица 2.3. Отметки лотков и глубина колодцев дворовой сети"),
    ("Табл. Б.2", "Справочно: табл. Б.2 СП 30.13330.2020 (NP → α)"),
]
for k, (sh, op) in enumerate(SODERZH):
    yach(w0, "A%d" % (4 + k), sh, f_bold)
    yach(w0, "B%d" % (4 + k), op)
yach(w0, "A17", "Обозначения: жёлтые ячейки — исходные данные (синий шрифт); зелёные — результаты; остальное — "
               "формулы. Табличные величины (K_s, h/d и V по Лукиным, напор насосов по кривым) внесены числами из "
               "расчёта и подписаны.", f_note, border=False, al=Alignment(horizontal="left", wrap_text=True))
w0.merge_cells("A17:B18")

# ---------------- Б.2 ----------------
wB = wb.create_sheet("Табл. Б.2")
yach(wB, "A1", "NP", f_bold, fill_h)
yach(wB, "B1", "α", f_bold, fill_h)
for k, (x, a) in enumerate(B2):
    wB.cell(row=2 + k, column=1, value=x).font = f_b
    wB.cell(row=2 + k, column=2, value=a).font = f_b
NAMES["B2_NP"] = "'Табл. Б.2'!$A$2:$A$%d" % (1 + len(B2))
NAMES["B2_A"] = "'Табл. Б.2'!$B$2:$B$%d" % (1 + len(B2))

# ---------------- 1.1 исходные ----------------
w = list_("1.1 Исходные", "Таблица 1.1. Исходные данные (вариант 3) и принятые допущения", [6, 62, 14, 12])
shapka(w, 3, ["№", "Показатель", "Значение", "Ед. изм."])
ISX = [  # имя, подпись, значение, ед., формат
    ("N_ET", "Этажность", IZ["N_ET"], "эт.", "0"),
    ("N_SEK", "Количество секций", IZ["N_SEK"], "шт.", "0"),
    ("H_ET", "Высота этажа", IZ["H_ET"], "м", "0.0"),
    ("U_ZAS", "Заселённость", REZ["U0"], "чел./кв.", "0.0"),
    ("H_PODVAL", "Высота подвала", IZ["H_PODVAL"], "м", "0.0"),
    ("DH_POL", "Высота пола 1-го этажа над планировкой", IZ["DH_POL"], "м", "0.0"),
    ("T_PEREKR", "Толщина перекрытий", IZ["T_PEREKR"], "м", "0.0"),
    ("H_PROMERZ", "Глубина промерзания грунта", IZ["H_PROMERZ"], "м", "0.0"),
    ("H_GAR", "Гарантийный напор в наружной сети", IZ["H_GAR"], "м", "0.0"),
    ("D_GOR_V", "Диаметр городского водопровода", IZ["D_GOR_V"], "мм", "0"),
    ("D_GOR_K", "Диаметр городской канализации", IZ["D_GOR_K"], "мм", "0"),
    ("NKV_SEK", "Квартир на этаже в секции (по плану — 3 санузла)", REZ["N_KV"], "шт.", "0"),
    ("N_PR", "Приборов в квартире (мойка, умывальник, ванна, унитаз)", 4, "шт.", "0"),
    ("K_L", "Коэффициент местных сопротивлений K_l (хоз.-питьевой, жилое здание)", 0.3, "—", "0.00"),
    ("H_SV", "Свободный напор у верхнего прибора (п. 8.21 СП 30)", 20.0, "м", "0.0"),
    ("Z_PR_DIKT", "Высота диктующего прибора над полом (душевая сетка над ванной)", 1.95, "м", "0.00"),
    ("Z_ZEML_GVK", "Отметка земли у ГВК (отн.)", -0.5, "м", "0.000"),
    ("L_VVOD", "Длина ввода: 16,5 + 6,5 до стены + 2,0 по подвалу", REZ["L_vvod"], "м", "0.0"),
    ("Z0_ABS", "Абсолютная отметка 0,000", REZ["Z0_ABS"], "м", "0.00"),
]
for k, (nm, txt, v, ed, fmt) in enumerate(ISX):
    r = 4 + k
    yach(w, "A%d" % r, k + 1)
    yach(w, "B%d" % r, txt)
    yach(w, "C%d" % r, v, f_in, fill_in, fmt)
    yach(w, "D%d" % r, ed)
    imya(nm, w, "C%d" % r)

# ---------------- 1.2 нормы ----------------
w = list_("1.2 Нормы", "Таблица 1.2. Нормы расхода воды (табл. А.2 СП 30.13330.2020, жилые дома с ваннами "
                       "1500–1700 мм с душами)", [62, 14, 14, 14])
shapka(w, 3, ["Норма", "общая (tot)", "горячей (h)", "холодной (c)"])
NR = [
    ("QU", "q_u — в сутки наибольшего водопотребления, л/сут на 1 жителя", NORMA["q_u_tot"], NORMA["q_u_h"], "minus"),
    ("QHR", "q_hr,u — в час наибольшего водопотребления, л/ч на 1 жителя", NORMA["q_hr_tot"], NORMA["q_hr_h"], "minus"),
    ("Q0", "q_0 — секундный расход прибором, л/с", NORMA["q0_tot"], NORMA["q0_c"], "same"),
    ("Q0HR", "q_0,hr — часовой расход прибором, л/ч", NORMA["q0hr_tot"], NORMA["q0hr_c"], "same"),
    ("T_CH", "T — время водопотребления, ч", NORMA["T"], NORMA["T"], "same"),
]
for k, (nm, txt, vt, vh, how) in enumerate(NR):
    r = 4 + k
    yach(w, "A%d" % r, txt)
    yach(w, "B%d" % r, vt, f_in, fill_in, "0.0##")
    yach(w, "C%d" % r, vh, f_in, fill_in, "0.0##")
    yach(w, "D%d" % r, "=B%d-C%d" % (r, r) if how == "minus" else "=C%d" % r, fmt="0.0##")
    imya(nm + "_TOT", w, "B%d" % r)
    imya(nm + "_H", w, "C%d" % r)
    imya(nm + "_C", w, "D%d" % r)
yach(w, "A10", "q_0 и q_0,hr холодной и горячей воды в табл. А.2 — одна графа; c = tot − h для суточной и часовой "
               "норм на жителя.", f_note, border=False)

# ---------------- 1.3 расходы ----------------
w = list_("1.3 Расходы", "Таблица 1.3. Расчётные расходы воды (п. 5 СП 30.13330.2020)", [58, 30, 15, 12])
shapka(w, 3, ["Величина", "Формула", "Значение", "Ед. изм."])
RS = [
    ("N_KV", "Квартир в здании n_кв", "n_эт·n_сек·n_кв.эт", "=N_ET*N_SEK*NKV_SEK", "шт.", "0"),
    ("U_", "Число жителей U (округление вверх)", "U₀·n_кв", "=ROUNDUP(U_ZAS*N_KV-1E-9,0)", "чел.", "0"),
    ("N_", "Число приборов N", "n_пр·n_кв", "=N_PR*N_KV", "шт.", "0"),
    ("QSUT_TOT", "Суточный общий Q_сут^tot", "q_u^tot·U/1000", "=QU_TOT*U_/1000", "м³/сут", "0.00"),
    ("QSUT_C", "Суточный холодной Q_сут^c", "q_u^c·U/1000", "=QU_C*U_/1000", "м³/сут", "0.00"),
    ("QT_TOT", "Среднечасовой общий q_T^tot", "Q_сут^tot/T", "=QSUT_TOT/T_CH_TOT", "м³/ч", "0.000"),
    ("QT_C", "Среднечасовой холодной q_T^c", "Q_сут^c/T", "=QSUT_C/T_CH_TOT", "м³/ч", "0.000"),
    ("P_TOT", "Вероятность действия P^tot", "q_hr,u^tot·U/(3600·q_0^tot·N)", "=QHR_TOT*U_/(3600*Q0_TOT*N_)", "—",
     "0.00000"),
    ("P_C", "Вероятность действия P^c", "q_hr,u^c·U/(3600·q_0^c·N)", "=QHR_C*U_/(3600*Q0_C*N_)", "—", "0.00000"),
    ("NP_TOT", "NP^tot", "N·P^tot", "=N_*P_TOT", "—", "0.000"),
    ("A_TOT", "α^tot (табл. Б.2, интерполяция)", "f(NP^tot)", "=" + alfa("NP_TOT"), "—", "0.000"),
    ("Q_TOT", "Максимальный секундный общий q^tot", "5·q_0^tot·α^tot", "=5*Q0_TOT*A_TOT", "л/с", "0.000"),
    ("NP_C", "NP^c", "N·P^c", "=N_*P_C", "—", "0.000"),
    ("A_C", "α^c", "f(NP^c)", "=" + alfa("NP_C"), "—", "0.000"),
    ("Q_C", "Максимальный секундный холодной q^c", "5·q_0^c·α^c", "=5*Q0_C*A_C", "л/с", "0.000"),
    ("PHR_TOT", "Часовая вероятность P_hr^tot", "3600·P^tot·q_0^tot/q_0,hr^tot", "=3600*P_TOT*Q0_TOT/Q0HR_TOT",
     "—", "0.00000"),
    ("AHR_TOT", "α_hr^tot", "f(N·P_hr^tot)", "=" + alfa("(N_*PHR_TOT)"), "—", "0.000"),
    ("QHR_TOT_", "Максимальный часовой общий q_hr^tot", "0,005·q_0,hr^tot·α_hr^tot", "=0.005*Q0HR_TOT*AHR_TOT",
     "м³/ч", "0.000"),
    ("PHR_C", "Часовая вероятность P_hr^c", "3600·P^c·q_0^c/q_0,hr^c", "=3600*P_C*Q0_C/Q0HR_C", "—", "0.00000"),
    ("AHR_C", "α_hr^c", "f(N·P_hr^c)", "=" + alfa("(N_*PHR_C)"), "—", "0.000"),
    ("QHR_C_", "Максимальный часовой холодной q_hr^c", "0,005·q_0,hr^c·α_hr^c", "=0.005*Q0HR_C*AHR_C", "м³/ч",
     "0.000"),
]
REZ_CELLS = {}
for k, (nm, txt, fo, f, ed, fmt) in enumerate(RS):
    r = 4 + k
    yach(w, "A%d" % r, txt)
    yach(w, "B%d" % r, fo)
    res = nm in ("U_", "N_", "QSUT_TOT", "QSUT_C", "Q_TOT", "Q_C", "QHR_TOT_", "QHR_C_")
    yach(w, "C%d" % r, f, f_bold if res else f_b, fill_res if res else None, fmt)
    yach(w, "D%d" % r, ed)
    imya(nm, w, "C%d" % r)
    REZ_CELLS[nm] = "'%s'!C%d" % (w.title, r)

# ---------------- 1.4 гидравлика ----------------
w = list_("1.4 Гидравлика В1", "Таблица 1.4. Гидравлический расчёт сети В1 по диктующему направлению "
                               "(стояк %s, душевая сетка 11-го этажа)" % REZ["st_dikt"],
          [8, 36, 9, 7, 11, 9, 9, 9, 7, 10, 9, 10, 9, 9, 9])
shapka(w, 3, ["Уча-сток", "Наименование", "L, м", "N", "P^c", "NP", "α", "q, л/с", "DN", "d_р, мм", "v, м/с",
              "i", "iL", "h, м", "Материал"])
PPR = {20: 13.2, 25: 16.6, 32: 21.2, 40: 26.6, 50: 33.2}      # ППР PN20: наружный -> внутренний, мм
VGP = {15: 15.7, 20: 21.2, 25: 27.1, 32: 35.9, 40: 41.0, 50: 53.0, 65: 67.5, 80: 80.5, 100: 105.0}
n_v1 = len(REZ["v1"])
for k, u in enumerate(REZ["v1"]):
    r = 4 + k
    uch = "%d–%s" % (u["n"], "НС" if u["n"] == n_v1 else str(u["n"] + 1))
    yach(w, "A%d" % r, uch)
    yach(w, "B%d" % r, u["imya"])
    yach(w, "C%d" % r, u["L"], f_in, fill_in, "0.00")
    yach(w, "D%d" % r, u["N"], f_in, fill_in, "0")
    yach(w, "E%d" % r, "=P_C", fmt="0.00000")
    yach(w, "F%d" % r, "=D%d*E%d" % (r, r), fmt="0.0000")
    yach(w, "G%d" % r, "=" + alfa("F%d" % r), fmt="0.000")
    yach(w, "H%d" % r, "=5*Q0_C*G%d" % r, fmt="0.000")
    yach(w, "I%d" % r, u["dn"], f_in, fill_in, "0")
    yach(w, "O%d" % r, "ППР" if u.get("mat") == "ппр" else "сталь", f_in, fill_in)
    yach(w, "J%d" % r, '=IF(O{r}="ППР",INDEX(PPR_D,MATCH(I{r},PPR_DN,0)),'
                       'INDEX(VGP_D,MATCH(I{r},VGP_DN,0))-1)'.format(r=r), fmt="0.0")
    yach(w, "K%d" % r, "=H%d/1000/(PI()*(J%d/1000)^2/4)" % (r, r), fmt="0.00")
    yach(w, "L%d" % r, '=IF(O{r}="ППР",0.000685*K{r}^1.774/(J{r}/1000)^1.226,'
                       'IF(K{r}<1.2,0.000912*K{r}^2/(J{r}/1000)^1.3*(1+0.867/K{r})^0.3,'
                       '0.00107*K{r}^2/(J{r}/1000)^1.3))'.format(r=r), fmt="0.0000")
    yach(w, "M%d" % r, "=L%d*C%d" % (r, r), fmt="0.000")
    yach(w, "N%d" % r, "=M%d*(1+K_L)" % r, fmt="0.000")
rs = 4 + n_v1
yach(w, "B%d" % rs, "Итого Σh по диктующему направлению", f_bold)
yach(w, "N%d" % rs, "=SUM(N4:N%d)" % (rs - 1), f_bold, fill_res, "0.00")
imya("SUM_H", w, "N%d" % rs)
REZ_CELLS["SUM_H"] = "'%s'!N%d" % (w.title, rs)
yach(w, "A%d" % (rs + 2), "Квартирные разводки, узлы учёта и стояки — ППР PN20 (ГОСТ 32415-2013): d_р — внутренний "
                          "диаметр, i = 0,000685·v^1,774/d_р^1,226 (Шевелёв, ф. 28). Подводки к стоякам, магистраль и "
                          "напорная линия — сталь ВГП (ГОСТ 3262-75): d_р — внутренний диаметр минус 1 мм, v < 1,2 м/с — "
                          "i = 0,000912·v²/d_р^1,3·(1 + 0,867/v)^0,3; v ≥ 1,2 м/с — i = 0,00107·v²/d_р^1,3. "
                          "h = iL·(1 + K_l). Диаметр — по экономичным скоростям (≤ 1,5 м/с в подводках, ≤ 1,2 м/с "
                          "в стояках и магистралях).", f_note, border=False,
     al=Alignment(horizontal="left", vertical="top", wrap_text=True))
w.merge_cells("A%d:O%d" % (rs + 2, rs + 3))
w.row_dimensions[rs + 2].height = 42
yach(w, "P3", "DN", f_bold, fill_h)
yach(w, "Q3", "d_вн, мм", f_bold, fill_h)
for k, (dn, d) in enumerate(sorted(VGP.items())):
    yach(w, "P%d" % (4 + k), dn, fmt="0")
    yach(w, "Q%d" % (4 + k), d, fmt="0.0")
NAMES["VGP_DN"] = "'%s'!$P$4:$P$%d" % (w.title, 3 + len(VGP))
NAMES["VGP_D"] = "'%s'!$Q$4:$Q$%d" % (w.title, 3 + len(VGP))
yach(w, "S3", "ППР Ø", f_bold, fill_h)
yach(w, "T3", "d_вн, мм", f_bold, fill_h)
for k, (dn, d) in enumerate(sorted(PPR.items())):
    yach(w, "S%d" % (4 + k), dn, fmt="0")
    yach(w, "T%d" % (4 + k), d, fmt="0.0")
NAMES["PPR_DN"] = "'%s'!$S$4:$S$%d" % (w.title, 3 + len(PPR))
NAMES["PPR_D"] = "'%s'!$T$4:$T$%d" % (w.title, 3 + len(PPR))
w.column_dimensions["R"].width = 3
w.column_dimensions["S"].width = 7
w.column_dimensions["T"].width = 10
w.column_dimensions["O"].width = 9
w.column_dimensions["P"].width = 7
w.column_dimensions["Q"].width = 10

# ---------------- 1.5 ввод ----------------
w = list_("1.5 Ввод", "Расчёт ввода В1-1 (ПЭ100 SDR17, потери по Колбруку–Уайту)", [52, 16, 12])
VV = [
    ("VV_Q", "Расчётный расход q^c", "=Q_C", "л/с", "0.000"),
    ("VV_DN", "Наружный диаметр ПЭ100 SDR17", REZ["vvod"]["dn"], "мм", "0"),
    ("VV_D", "Внутренний диаметр", 55.4, "мм", "0.0"),
    ("VV_DELTA", "Эквивалентная шероховатость Δ", 0.01, "мм", "0.00"),
    ("VV_NU", "Кинематическая вязкость воды (10 °С)", 1.31E-6, "м²/с", "0.00E+00"),
    ("VV_V", "Скорость v = q/(π d²/4)", "=VV_Q/1000/(PI()*(VV_D/1000)^2/4)", "м/с", "0.000"),
    ("VV_RE", "Число Рейнольдса Re = v·d/ν", "=VV_V*VV_D/1000/VV_NU", "—", "0"),
]
shapka(w, 3, ["Величина", "Значение", "Ед. изм."])
r = 4
for nm, txt, v, ed, fmt in VV:
    yach(w, "A%d" % r, txt)
    inp = not (isinstance(v, str) and v.startswith("="))
    yach(w, "B%d" % r, v, f_in if inp else f_b, fill_in if inp else None, fmt)
    yach(w, "C%d" % r, ed)
    imya(nm, w, "B%d" % r)
    r += 1
yach(w, "A%d" % r, "λ — итерации формулы Колбрука–Уайта: 1/√λ = −2·lg(Δ/(3,7d) + 2,51/(Re·√λ))", f_bold)
r += 1
prev = None
for k in range(8):
    yach(w, "A%d" % r, "λ, итерация %d" % (k + 1))
    lam_prev = "0.02" if prev is None else "B%d" % prev
    yach(w, "B%d" % r, "=(-2*LOG10(VV_DELTA/1000/(3.7*VV_D/1000)+2.51/(VV_RE*SQRT(%s))))^-2" % lam_prev,
         fmt="0.00000")
    prev = r
    r += 1
imya("VV_LAM", w, "B%d" % prev)
for nm, txt, f, ed, fmt in [
        ("VV_I", "Удельные потери i = λ/d·v²/(2g)", "=VV_LAM/(VV_D/1000)*VV_V^2/(2*9.81)", "м/м", "0.00000"),
        ("H_VVOD", "Потери на вводе h_ввод = i·L·(1 + K_l)", "=VV_I*L_VVOD*(1+K_L)", "м", "0.000")]:
    yach(w, "A%d" % r, txt, f_bold if nm == "H_VVOD" else f_b)
    yach(w, "B%d" % r, f, f_bold if nm == "H_VVOD" else f_b, fill_res if nm == "H_VVOD" else None, fmt)
    yach(w, "C%d" % r, ed)
    imya(nm, w, "B%d" % r)
    REZ_CELLS[nm] = "'%s'!B%d" % (w.title, r)
    r += 1

# ---------------- 1.6 счётчики ----------------
w = list_("1.6 Счётчики", "Таблица 1.5. Подбор счётчиков воды (табл. 12.1, п. 12.15–12.16 СП 30.13330.2020)",
          [26, 8, 16, 14, 18, 12, 12, 13, 14])
shapka(w, 3, ["Узел", "DN", "q_экспл, м³/ч", "S, м/(л/с)²", "Расчётный расход q, л/с", "h_сч = S·q², м",
              "Допустимо, м", "q_T ≤ q_экспл", "Вывод"])
sd, sk = REZ["schetchik_dom"], REZ["sch_kv"]
yach(w, "A4", "Общедомовой ВСХН-25 (ввод)")
yach(w, "B4", sd[0], f_in, fill_in, "0")
yach(w, "C4", sd[1], f_in, fill_in, "0.0")
yach(w, "D4", sd[2], f_in, fill_in, "0.00")
yach(w, "E4", "=Q_C", fmt="0.000")
yach(w, "H4", '=IF(QT_C<=C4,"да","нет")')
yach(w, "A5", "Квартирный (4 прибора)")
yach(w, "B5", sk[0], f_in, fill_in, "0")
yach(w, "C5", sk[1], f_in, fill_in, "0.0")
yach(w, "D5", sk[2], f_in, fill_in, "0.0")
yach(w, "E5", "=5*Q0_C*" + alfa("(N_PR*P_C)"), fmt="0.000")
yach(w, "H5", "—")
for r in (4, 5):
    yach(w, "F%d" % r, "=D%d*E%d^2" % (r, r), f_bold, fill_res, "0.00")
    yach(w, "G%d" % r, 5.0, f_in, fill_in, "0.0")
    yach(w, "I%d" % r, '=IF(F{r}<=G{r},"подходит","больше DN")'.format(r=r))
yach(w, "A6", "Σh_сч (общедомовой + квартирный)", f_bold)
yach(w, "F6", "=F4+F5", f_bold, fill_res, "0.00")
imya("H_SCH", w, "F6")
REZ_CELLS["H_SCH"] = "'%s'!F6" % w.title

# ---------------- 1.7 напор ----------------
w = list_("1.7 Напор", "Требуемый напор и напор повысительной установки", [62, 34, 14, 10])
shapka(w, 3, ["Величина", "Формула", "Значение", "Ед. изм."])
NP_ = [
    ("Z_GOR", "Верх городской трубы (отн.)", "z_земли − (H_пр + 0,5)", "=Z_ZEML_GVK-(H_PROMERZ+0.5)", "м", "0.000"),
    ("Z_DIKT", "Отметка диктующего прибора (11-й этаж)", "(n_эт − 1)·h_эт + 1,95", "=(N_ET-1)*H_ET+Z_PR_DIKT", "м",
     "0.000"),
    ("H_GEOM", "Геометрическая высота H_geom", "z_дикт − z_гор", "=Z_DIKT-Z_GOR", "м", "0.00"),
    ("H_SUMH", "Потери в сети Σh (табл. 1.4)", "Σh", "=SUM_H", "м", "0.00"),
    ("H_SV_", "Свободный напор H_f", "п. 8.21", "=H_SV", "м", "0.0"),
    ("H_SCH_", "Потери в счётчиках Σh_сч", "табл. 1.5", "=H_SCH", "м", "0.00"),
    ("H_VVOD_", "Потери на вводе h_ввод", "лист 1.5", "=H_VVOD", "м", "0.00"),
    ("H_TR", "Требуемый напор H_тр", "H_geom + Σh + H_f + Σh_сч + h_ввод",
     "=H_GEOM+H_SUMH+H_SV_+H_SCH_+H_VVOD_", "м", "0.00"),
    ("H_RAB", "Напор установки H_р", "H_тр − H_гар", "=H_TR-H_GAR", "м", "0.00"),
    ("Q_NAS", "Подача установки Q", "q^c·3,6", "=Q_C*3.6", "м³/ч", "0.00"),
    ("N_POL", "Полезная мощность N", "ρ·g·Q·H_р", "=1000*9.81*Q_NAS/3600*H_RAB/1000", "кВт", "0.00"),
]
for k, (nm, txt, fo, f, ed, fmt) in enumerate(NP_):
    r = 4 + k
    yach(w, "A%d" % r, txt, f_bold if nm in ("H_TR", "H_RAB") else f_b)
    yach(w, "B%d" % r, fo)
    yach(w, "C%d" % r, f, f_bold if nm in ("H_TR", "H_RAB") else f_b,
         fill_res if nm in ("H_TR", "H_RAB") else None, fmt)
    yach(w, "D%d" % r, ed)
    imya(nm, w, "C%d" % r)
    REZ_CELLS[nm] = "'%s'!C%d" % (w.title, r)
r0 = 4 + len(NP_) + 1
yach(w, "A%d" % r0, "Регуляторы давления (п. 8.22 СП 30: напор у нижнего прибора более 45 м)", f_bold, border=False)
shapka(w, r0 + 1, ["Этаж", "Напор у смесителя ванны (0,8 м) без расхода, м", "Регулятор", ""])
for k in range(IZ["N_ET"]):
    r = r0 + 2 + k
    yach(w, "A%d" % r, k + 1, al=CEN)
    yach(w, "B%d" % r, "=H_GAR+H_RAB-((A{r}-1)*H_ET+0.8-Z_GOR)".format(r=r), fmt="0.0")
    yach(w, "C%d" % r, '=IF(B{r}>45,"нужен","нет")'.format(r=r))
REZ_CELLS["ET1"] = "'%s'!B%d" % (w.title, r0 + 2)

# ---------------- 1.8 насос ----------------
w = list_("1.8 Насос", "Таблица 1.6. Подбор насоса (Wilo MVL 4, напор при 50 Гц по кривым каталога при расчётной "
                       "подаче)", [14, 18, 13, 16, 11, 13, 16])
shapka(w, 3, ["Насос", "H при 50 Гц, м", "Запас, м", "В рабочей зоне", "КПД, %", "Мотор, кВт", "Вывод"])
kand = []
for line in open(os.path.join(KURS, "PODBOR_NASOSA.md"), encoding="utf-8"):
    m = re.match(r"\| (MVL \d+) \| ([^|]+) \| ([^|]+) \| ([^|]+) \| ([^|]+) \| ([^|]+) \| ([^|]+) \|", line)
    if m:
        kand.append([c.strip().replace("**", "") for c in m.groups()])
for k, c in enumerate(kand):
    r = 4 + k
    yach(w, "A%d" % r, c[0])
    yach(w, "B%d" % r, float(c[1].replace(",", ".")), f_in, fill_in, "0.0")
    yach(w, "C%d" % r, "=B%d-H_RAB" % r, fmt="0.0")
    yach(w, "D%d" % r, c[3])
    yach(w, "E%d" % r, float(c[4].replace(",", ".")), f_in, fill_in, "0.0")
    yach(w, "F%d" % r, float(c[6].replace(",", ".")), f_in, fill_in, "0.0")
    r1, r2 = 4, 3 + len(kand)
    yach(w, "G%d" % r, '=IF(C{r}<0,"мал напор",IF(B{r}=MINIFS_,"принят","подходит"))'.format(r=r).replace(
        "MINIFS_", "SUMPRODUCT(MIN(B{a}:B{b}+(C{a}:C{b}<0)*1000))".format(a=r1, b=r2)))
yach(w, "A%d" % (5 + len(kand)), "Принята установка Wilo %s: 1 рабочий + 1 резервный насос, частотное "
                                 "регулирование. Напор насосов — по кривым каталога WILO серии MVL при "
                                 "Q = %s м³/ч." % (NAS["tip"], str(NAS["Q"]).replace(".", ",")), f_note, border=False)

# ---------------- 2.1 стояки ----------------
w = list_("2.1 Стояки К1", "Таблица 2.1. Пропускная способность стояков К1 (полипропилен Ø110, отводы 45°)",
          [12, 8, 12, 10, 13, 13, 15, 14])
shapka(w, 3, ["Стояк", "N", "NP^tot", "α", "q^tot, л/с", "q^s, л/с", "Допустимо, л/с", "Вывод"])
for k, s_ in enumerate(REZ["k1_stoyaki"]):
    r = 4 + k
    yach(w, "A%d" % r, s_["st"])
    yach(w, "B%d" % r, "=N_PR*N_ET", fmt="0")
    yach(w, "C%d" % r, "=B%d*P_TOT" % r, fmt="0.0000")
    yach(w, "D%d" % r, "=" + alfa("C%d" % r), fmt="0.000")
    yach(w, "E%d" % r, "=5*Q0_TOT*D%d" % r, fmt="0.000")
    yach(w, "F%d" % r, "=E%d+1.6" % r, f_bold, fill_res, "0.000")
    yach(w, "G%d" % r, s_["q_dop"], f_in, fill_in, "0.00")
    yach(w, "H%d" % r, '=IF(F{r}<=G{r},"обеспечена","нет")'.format(r=r))
    if k == 0:
        REZ_CELLS["QS_ST"] = "'%s'!F%d" % (w.title, r)
yach(w, "A11", "q^s = q^tot + q_0^s, q_0^s = 1,6 л/с (унитаз со смывным бачком, табл. А.1). Допустимый расход — прил. К "
               "СП 30.13330.2020.", f_note, border=False)

# ---------------- 2.2 горизонтальные ----------------
w = list_("2.2 Горизонт. К1", "Таблица 2.2. Расчёт горизонтальных участков К1 и дворовой сети",
          [34, 8, 7, 10, 9, 11, 9, 11, 9, 8, 8, 9, 11, 9])
shapka(w, 3, ["Участок", "L, м", "N", "P_hr^tot", "α_hr", "q_hr, м³/ч", "K_s", "q^sL, л/с", "d, мм", "i", "h/d",
              "V, м/с", "V·√(h/d)", "iL, м"])
for k, u in enumerate(REZ["k1"]):
    r = 4 + k
    yach(w, "A%d" % r, u["imya"])
    yach(w, "B%d" % r, u["L"], f_in, fill_in, "0.00")
    yach(w, "C%d" % r, u["N"], f_in, fill_in, "0")
    yach(w, "D%d" % r, "=PHR_TOT", fmt="0.00000")
    yach(w, "E%d" % r, "=" + alfa("(C%d*D%d)" % (r, r)), fmt="0.000")
    yach(w, "F%d" % r, "=0.005*Q0HR_TOT*E%d" % r, fmt="0.000")
    yach(w, "G%d" % r, round(u["Ks"], 4), f_in, fill_in, "0.00")
    yach(w, "H%d" % r, "=F%d/3.6+G%d*1.1" % (r, r), f_bold, fill_res, "0.00")
    yach(w, "I%d" % r, u["d"], f_in, fill_in, "0")
    yach(w, "J%d" % r, u["i"], f_in, fill_in, "0.000")
    yach(w, "K%d" % r, round(u["hd"], 3), f_in, fill_in, "0.00")
    yach(w, "L%d" % r, round(u["v"], 3), f_in, fill_in, "0.00")
    yach(w, "M%d" % r, "=L%d*SQRT(K%d)" % (r, r), fmt="0.00")
    yach(w, "N%d" % r, "=B%d*J%d" % (r, r), fmt="0.000")
    imya("K1_IL_%d" % k, w, "N%d" % r)
    imya("K1_I_%d" % k, w, "J%d" % r)
n_k1 = len(REZ["k1"])
yach(w, "A%d" % (5 + n_k1), "q^sL = q_hr/3,6 + K_s·q_0^s1 (q_0^s1 = 1,1 л/с — ванна). K_s — по табл. 5.1 СП 30 с "
                            "интерполяцией по L и N; h/d и V — по формуле Н.Н. Павловского (n = 0,014, таблицы "
                            "Лукиных) при принятых d и i. Все участки безрасчётные: условие V·√(h/d) ≥ K невыполнимо "
                            "при малых расходах, уклон i ≥ 1/d.", f_note, border=False,
     al=Alignment(horizontal="left", vertical="top", wrap_text=True))
w.merge_cells("A%d:N%d" % (5 + n_k1, 6 + n_k1))
w.row_dimensions[5 + n_k1].height = 28

# ---------------- 2.3 отметки ----------------
w = list_("2.3 Отметки", "Таблица 2.3. Отметки лотков и глубина колодцев дворовой сети К1 (абсолютные), м",
          [30, 12, 12, 14, 14, 30])
shapka(w, 3, ["Точка", "Земля", "Лоток", "Глубина лотка", "Не мельче H₀", "Как получен лоток"])
zem = REZ["zeml"]
# индексы участков дворовой сети в REZ["k1"]
idx = {u["imya"].split(" (")[0]: k for k, u in enumerate(REZ["k1"])}
vyp = next(k for k, u in enumerate(REZ["k1"]) if u["imya"].startswith("выпуск"))
OTM = [
    ("Выпуск у стены (−2,000)", None, "=Z0_ABS-2", "лоток выпуска у стены"),
    ("Выпуск в КК-1 (Ø100)", None, "=C4-K1_I_%d*5" % vyp, "−i·5 м (стена — колодец)"),
    ("КК-1", zem["КК-1"], "=C5+0.1-0.15", "шелыга в шелыгу: Ø100 → Ø150"),
    ("КК-3", zem["КК-3"], "=C6-K1_IL_%d" % idx["КК-1 - КК-3"], "− i·L участка КК-1–КК-3 (КК-2 убран)"),
    ("ККК", zem["ККК"], "=C7-K1_IL_%d" % idx["КК-3 - ККК"], "− i·L участка КК-3–ККК"),
    ("ГКК", zem["ГКК"], "=C8-K1_IL_%d" % idx["ККК - ГКК"], "− i·L участка ККК–ГКК"),
]
for k, (nm, z, f, kak) in enumerate(OTM):
    r = 4 + k
    yach(w, "A%d" % r, nm)
    yach(w, "B%d" % r, z if z is not None else "—", f_in if z is not None else f_b, fill_in if z is not None else None,
         "0.00")
    yach(w, "C%d" % r, f, fmt="0.00")
    if z is not None:
        yach(w, "D%d" % r, "=B%d-C%d" % (r, r), f_bold, fill_res, "0.00")
        yach(w, "E%d" % r, '=IF(D{r}>=H0_MIN,"да","нет")'.format(r=r))
    else:
        yach(w, "D%d" % r, "—")
        yach(w, "E%d" % r, "—")
    yach(w, "F%d" % r, kak)
    if nm == "ГКК":
        REZ_CELLS["LOT_GKK"] = "'%s'!C%d" % (w.title, r)
        r_gkk = r
yach(w, "A12", "H₀ = max(H_пр − 0,3; 0,7 + d)", f_bold)
yach(w, "C12", "=MAX(H_PROMERZ-0.3,0.7+0.15)", f_bold, fill_res, "0.00")
imya("H0_MIN", w, "C12")
yach(w, "A13", "Лоток городского коллектора Ø400 в ГКК («по месту»)")
yach(w, "C13", REZ["Z_GKK_GOR"], f_in, fill_in, "0.00")
yach(w, "A14", "Дворовая труба входит выше шелыги коллектора на, м")
yach(w, "C14", "=C%d-(C13+0.4)" % r_gkk, fmt="0.00")    # лоток дворовой трубы над шелыгой коллектора
REZ_CELLS["PEREPAD"] = "'%s'!C14" % w.title

# справочная таблица — в конец книги
wb.move_sheet("Табл. Б.2", offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("Табл. Б.2"))

# ---------------- имена ----------------
for nm, ref in NAMES.items():
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = ws.title in ("Табл. Б.2",)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
json.dump(REZ_CELLS, open(os.path.join(HERE, "_yacheyki_proverki.json"), "w", encoding="utf-8"), ensure_ascii=False)
print(OUT)
