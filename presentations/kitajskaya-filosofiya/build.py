# -*- coding: utf-8 -*-
"""Сборка презентации «Китайская философия: даосизм и конфуцианство».

Стиль — новый китайский минимализм: рисовая бумага, тушь, выдержанная палитра.
Шрифт — Arial (по требованию). На слайде только выжимка, текст выступления —
в заметках докладчика. Результат полностью редактируемый .pptx.

    python3 build.py
"""
import os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

from PIL import ImageFont

import content

ROOT = os.path.dirname(os.path.abspath(__file__))
BG = os.path.join(ROOT, "assets", "bg")
IMG = os.path.join(ROOT, "assets", "img")
OUT = os.path.join(ROOT, "out")

# ── холст ───────────────────────────────────────────────────────────────────
SW, SH = 13.333, 7.5          # дюймы, 16:9
M = 0.95                      # поле

# ── палитра (выдержанная) ───────────────────────────────────────────────────
PAPER      = RGBColor(0xED, 0xE9, 0xE1)
PAPER_SOFT = RGBColor(0xE3, 0xDE, 0xD2)
INK        = RGBColor(0x2A, 0x2A, 0x24)
INK_SOFT   = RGBColor(0x4A, 0x49, 0x42)
GREEN      = RGBColor(0x2C, 0x38, 0x2F)
GREEN_SOFT = RGBColor(0x6E, 0x7D, 0x6C)
MUTED      = RGBColor(0x8A, 0x85, 0x78)
HAIR       = RGBColor(0xBE, 0xB8, 0xAA)
BRONZE     = RGBColor(0x8C, 0x7B, 0x5F)

# ── шрифты ──────────────────────────────────────────────────────────────────
FONT = "Arial"                # требование заказчика
EA   = "Microsoft YaHei"      # подстановка для иероглифов


LIB = "/usr/share/fonts/truetype/liberation/LiberationSans-%s.ttf"
_FONTS = {}


def _metric_font(size, bold=False, italic=False):
    """Liberation Sans метрически совпадает с Arial — годится для расчётов."""
    style = ("Bold" if bold else "") + ("Italic" if italic else "")
    style = style or "Regular"
    if style == "BoldItalic":
        style = "BoldItalic"
    key = (round(size * 4), style)
    if key not in _FONTS:
        _FONTS[key] = ImageFont.truetype(LIB % style, int(round(size * 4)))
    return _FONTS[key]


def n_lines(text, size, width, bold=False, italic=False):
    """Сколько строк займёт текст в рамке шириной width дюймов."""
    f = _metric_font(size, bold, italic)
    limit = width * 72 * 4
    total = 0
    for src in str(text).split("\n"):
        cur, lines = "", 1
        for word in src.split():
            probe = (cur + " " + word).strip()
            if f.getlength(probe) <= limit:
                cur = probe
            else:
                lines += 1
                cur = word
        total += lines
    return total


def text_h(text, size, width, line=1.3, bold=False, italic=False):
    """Высота текстового блока в дюймах."""
    return n_lines(text, size, width, bold, italic) * size * line / 72.0


# ════════════════════════════════════════════════════════════════ примитивы ══
def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def para(tf, text, size, *, bold=False, color=INK, align=PP_ALIGN.LEFT,
         spc=0, line=1.25, before=0, after=0, italic=False, first=False,
         font=FONT, caps=False):
    """Абзац. spc — разрядка в пунктах, line — межстрочный множитель."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = line
    if before:
        p.space_before = Pt(before)
    if after:
        p.space_after = Pt(after)
    r = p.add_run()
    r.text = text.upper() if caps else text
    f = r.font
    f.name = font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    # восточноазиатская подстановка для иероглифов
    rPr = r._r.get_or_add_rPr()
    ea = rPr.makeelement(qn("a:ea"), {"typeface": EA})
    rPr.append(ea)
    if spc:
        rPr.set("spc", str(int(spc * 100)))
    return p


def block(tf, lines, size, **kw):
    """Многострочный текст: каждая строка — отдельный абзац."""
    out = []
    for i, ln in enumerate(str(lines).split("\n")):
        out.append(para(tf, ln, size, first=(i == 0), **kw))
    return out


def rect(slide, x, y, w, h, color, transparency=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    if transparency:
        sp = sh.fill.fore_color._xFill.find(qn("a:srgbClr"))
        alpha = sp.makeelement(qn("a:alpha"), {"val": str(int((1 - transparency) * 100000))})
        sp.append(alpha)
    return sh


def rule(slide, x, y, w, color=HAIR, h=0.014):
    return rect(slide, x, y, w, h, color)


def vrule(slide, x, y, h, color=HAIR, w=0.014):
    return rect(slide, x, y, w, h, color)


def find(name):
    """Ищет иллюстрацию, не привязываясь к расширению."""
    p = os.path.join(IMG, name)
    if os.path.exists(p):
        return p
    stem = os.path.splitext(name)[0]
    for ext in (".png", ".jpg", ".jpeg"):
        p = os.path.join(IMG, stem + ext)
        if os.path.exists(p):
            return p
    raise FileNotFoundError(name)


def pic_cover(slide, name, x, y, w, h):
    """Картинка заполняет рамку целиком, лишнее обрезается."""
    path = find(name)
    iw, ih = Image.open(path).size
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    src, tgt = iw / ih, w / h
    if src > tgt:
        k = (1 - tgt / src) / 2
        pic.crop_left = pic.crop_right = k
    else:
        k = (1 - src / tgt) / 2
        pic.crop_top = pic.crop_bottom = k
    return pic


def pic_fit(slide, name, x, y, w, h, align="center"):
    """Картинка вписывается в рамку целиком, пропорции сохраняются."""
    path = find(name)
    iw, ih = Image.open(path).size
    k = min(w / iw, h / ih)
    pw, ph = iw * k, ih * k
    px = x + (w - pw) / 2 if align == "center" else x
    py = y + (h - ph) / 2
    return slide.shapes.add_picture(path, Inches(px), Inches(py),
                                    Inches(pw), Inches(ph))


def background(slide, name):
    slide.shapes.add_picture(os.path.join(BG, name), 0, 0,
                             Inches(SW), Inches(SH))


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text or ""


def footer(slide, num, label="", light=False):
    col = RGBColor(0xB9, 0xC2, 0xB6) if light else MUTED
    tf = textbox(slide, M, SH - 0.62, 7.0, 0.3)
    para(tf, label, 9, color=col, spc=2.2, first=True, caps=True)
    tf = textbox(slide, SW - M - 1.2, SH - 0.62, 1.2, 0.3)
    para(tf, "%02d" % num, 9.5, color=col, align=PP_ALIGN.RIGHT, first=True)


def head(slide, kicker, title, *, y=0.78, size=30, color=INK, width=9.2):
    """Стандартная шапка слайда: рубрика + заголовок + линейка."""
    if kicker:
        tf = textbox(slide, M, y, width, 0.3)
        para(tf, kicker, 10, color=MUTED, spc=3.0, first=True, caps=True)
        y += 0.42
    lines = str(title).split("\n")
    h = len(lines) * size * 1.10 / 72.0
    tf = textbox(slide, M, y, width, h + 0.15)
    for i, ln in enumerate(lines):
        para(tf, ln, size, bold=True, color=color, spc=0.8, line=1.08,
             first=(i == 0))
    y += h + 0.24
    rule(slide, M, y, 1.55, color=BRONZE)
    return y + 0.42


def bullets(slide, items, x, y, w, *, size=16, line=1.32, gap=0.26,
            color=INK, dash=True, indent=0.44):
    """Список тезисов с точным вертикальным ритмом."""
    tw = w - indent
    for it in items:
        h = text_h(it, size, tw, line)
        if dash:
            tf = textbox(slide, x, y + 0.02, 0.32, 0.32)
            para(tf, "—", size - 1, color=BRONZE, first=True)
        tf = textbox(slide, x + indent, y, tw, h + 0.08)
        para(tf, it, size, color=color, line=line, first=True)
        y += h + gap
    return y


def lead_line(slide, text, x, y, w, size=16, color=GREEN):
    h = text_h(text, size, w, 1.34, italic=True)
    tf = textbox(slide, x, y, w, h + 0.08)
    para(tf, text, size, color=color, italic=True, line=1.34, first=True)
    return y + h + 0.30


def highlight(slide, text, size=14.5):
    """Акцентная плашка внизу слайда."""
    w = SW - 2 * M - 0.34
    h = text_h(text, size, w, 1.34)
    top = SH - 0.92 - h
    rule(slide, M, top - 0.22, SW - 2 * M, color=HAIR, h=0.009)
    vrule(slide, M, top, h + 0.04, color=GREEN, w=0.028)
    tf = textbox(slide, M + 0.30, top, w, h + 0.1)
    para(tf, text, size, color=GREEN, line=1.34, first=True)
    return top


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


# ═══════════════════════════════════════════════════════════════════ макеты ══
def s_cover(prs, d, num):
    s = blank(prs)
    background(s, "cover.png")

    band = 3.58
    rect(s, 0, band, SW, SH - band, PAPER, transparency=0.04)
    rule(s, 0, band, SW, color=HAIR, h=0.012)

    y = band + 0.52
    tf = textbox(s, M, y, 6.2, 0.3)
    para(tf, d["kicker"], 10.5, color=MUTED, spc=3.4, first=True)

    tf = textbox(s, M, y + 0.44, 6.6, 1.6)
    for i, ln in enumerate(d["title"].split("\n")):
        para(tf, ln, 40, bold=True, color=INK, spc=1.2, line=1.06, first=(i == 0))

    rule(s, M, y + 1.92, 2.0, color=BRONZE)
    tf = textbox(s, M, y + 2.14, 6.0, 0.4)
    para(tf, d["subtitle"], 19, color=INK_SOFT, spc=0.6, first=True)

    qh = text_h("«%s»" % d["quote"], 14.5, 4.3, 1.55, italic=True)
    vrule(s, 7.62, y + 0.08, qh + 0.62, color=HAIR)
    tf = textbox(s, 7.95, y + 0.08, 4.3, qh + 0.2)
    para(tf, "«%s»" % d["quote"], 14.5, color=INK_SOFT, line=1.55, italic=True,
         first=True)
    tf = textbox(s, 7.95, y + 0.42 + qh, 4.3, 0.3)
    para(tf, d["quote_author"], 10, color=MUTED, spc=2.4, first=True, caps=True)
    tf = textbox(s, 7.95, y + 2.14, 4.3, 0.3)
    para(tf, d["meta"], 11, color=MUTED, spc=2.0, first=True)

    tf = textbox(s, SW - 1.5, 0.58, 0.9, 2.2)
    para(tf, "道", 28, color=RGBColor(0x55, 0x57, 0x50), align=PP_ALIGN.CENTER, first=True)
    para(tf, "德", 28, color=RGBColor(0x55, 0x57, 0x50), align=PP_ALIGN.CENTER, before=6)
    notes(s, d["notes"])
    return s


def s_contents(prs, d, num):
    s = blank(prs)
    background(s, "paper-clean.png")
    px = 9.42
    s.shapes.add_picture(os.path.join(BG, "green-panel.png"), Inches(px), 0,
                         Inches(SW - px), Inches(SH))

    tf = textbox(s, px + 0.55, 1.05, 2.6, 1.2)
    para(tf, d["cjk"], 40, color=PAPER, first=True)
    tf = textbox(s, px + 0.58, 2.34, 2.8, 0.4)
    para(tf, d["title"], 13, color=RGBColor(0xC2, 0xCB, 0xBF), spc=3.2, first=True)
    rule(s, px + 0.58, 2.92, 0.9, color=RGBColor(0x8F, 0x9C, 0x8C))

    y, step = 0.96, 0.86
    for i, (n, t, sub) in enumerate(d["items"]):
        tf = textbox(s, M, y + 0.04, 0.8, 0.4)
        para(tf, n, 15, color=BRONZE, spc=1.0, first=True)
        tf = textbox(s, M + 0.85, y, 6.6, 0.42)
        para(tf, t, 19, bold=True, color=INK, spc=0.4, first=True)
        tf = textbox(s, M + 0.85, y + 0.38, 6.6, 0.3)
        para(tf, sub, 12.5, color=MUTED, first=True)
        if i < len(d["items"]) - 1:
            rule(s, M, y + step - 0.14, 7.6, color=HAIR, h=0.008)
        y += step
    notes(s, d["notes"])
    footer(s, num, "Содержание")
    return s


def s_section(prs, d, num):
    s = blank(prs)
    background(s, "green-full.png")

    tf = textbox(s, SW - 4.6, 1.15, 3.6, 3.2)
    para(tf, d["cjk"], 128, color=RGBColor(0x3B, 0x49, 0x3D),
         align=PP_ALIGN.RIGHT, first=True)

    tf = textbox(s, M, 2.28, 7.0, 0.3)
    para(tf, d["number"], 11, color=RGBColor(0x9A, 0xA6, 0x97), spc=3.4, first=True)

    tf = textbox(s, M, 2.72, 8.4, 1.1)
    para(tf, d["title"], 46, bold=True, color=PAPER, spc=2.0, first=True)

    rule(s, M, 3.98, 1.7, color=RGBColor(0x8C, 0x7B, 0x5F))

    tf = textbox(s, M, 4.22, 7.0, 0.4)
    para(tf, d["subtitle"], 17, color=RGBColor(0xC2, 0xCB, 0xBF), spc=0.6, first=True)

    cw = (SW - 2 * M) / max(len(d["items"]), 1)
    x = M
    for it in d["items"]:
        tf = textbox(s, x, 5.34, cw - 0.35, 0.3)
        para(tf, "—", 12, color=RGBColor(0x8C, 0x7B, 0x5F), first=True)
        tf = textbox(s, x, 5.62, cw - 0.35, 0.7)
        para(tf, it, 13.5, color=RGBColor(0xD3, 0xD8, 0xCF), line=1.3, first=True)
        x += cw
    notes(s, d["notes"])
    footer(s, num, d["title"], light=True)
    return s


def _image_panel(slide, d, x=8.0, top=1.28, bottom=5.48):
    """Иллюстрация справа + подпись."""
    w = SW - M - x
    h = bottom - top
    if d.get("img_mode") == "fit":
        pic_fit(slide, d["img"], x, top, w, h)
    else:
        pic_cover(slide, d["img"], x, top, w, h)
    if d.get("caption"):
        ch = text_h(d["caption"], 10.5, w, 1.3)
        tf = textbox(slide, x, bottom + 0.16, w, ch + 0.1)
        al = PP_ALIGN.CENTER if d.get("img_mode") == "fit" else PP_ALIGN.LEFT
        para(tf, d["caption"], 10.5, color=MUTED, line=1.3, align=al, first=True)


def s_text_image(prs, d, num):
    s = blank(prs)
    background(s, "paper-clean.png")
    y = head(s, d["kicker"], d["title"], width=7.0)
    col = 6.6

    if d.get("lead"):
        y = lead_line(s, d["lead"], M, y, col)
    bullets(s, d["points"], M, y, col, size=16)

    _image_panel(s, d)
    if d.get("highlight"):
        highlight(s, d["highlight"])
    notes(s, d["notes"])
    footer(s, num, d["kicker"])
    return s


def s_points(prs, d, num):
    s = blank(prs)
    background(s, "paper-clean.png")
    y = head(s, d["kicker"], d["title"], width=7.4)
    col = 6.6

    if d.get("lead"):
        y = lead_line(s, d["lead"], M, y, col)
    bullets(s, d["points"], M, y, col, size=16.5)

    if d.get("img"):
        _image_panel(s, d, top=1.28, bottom=5.60)
    if d.get("highlight"):
        highlight(s, d["highlight"])
    notes(s, d["notes"])
    footer(s, num, d["kicker"])
    return s


def s_two_portraits(prs, d, num):
    s = blank(prs)
    background(s, "paper-clean.png")
    head(s, d["kicker"], d["title"], width=8.0)

    top, ph = 2.15, 2.85
    vrule(s, SW / 2 - 0.01, top - 0.10, ph + 0.30, color=HAIR)
    for side, x in ((d["left"], M + 0.15), (d["right"], 7.30)):
        pic_fit(s, side["img"], x, top, 1.95, ph, align="left")
        tx = x + 2.30
        tw = 2.95
        tf = textbox(s, tx, top + 0.06, tw, 0.5)
        para(tf, side["name"], 22, bold=True, color=INK, spc=1.4, first=True)
        tf = textbox(s, tx, top + 0.62, tw, 0.3)
        para(tf, side["dates"], 12.5, color=BRONZE, spc=1.4, first=True)
        rule(s, tx, top + 1.10, 0.8, color=HAIR)
        dh = text_h(side["desc"], 15, tw, 1.34)
        tf = textbox(s, tx, top + 1.32, tw, dh + 0.1)
        block(tf, side["desc"], 15, color=INK_SOFT, line=1.34)

    if d.get("footnote"):
        fh = text_h(d["footnote"], 14, SW - 2 * M, 1.3, italic=True)
        top2 = SH - 0.92 - fh
        rule(s, M, top2 - 0.24, SW - 2 * M, color=HAIR, h=0.009)
        tf = textbox(s, M, top2, SW - 2 * M, fh + 0.1)
        para(tf, d["footnote"], 14, color=MUTED, italic=True, first=True)
    notes(s, d["notes"])
    footer(s, num, d["kicker"])
    return s


def s_compare(prs, d, num):
    s = blank(prs)
    background(s, "paper-clean.png")
    top = head(s, d["kicker"], d["title"], width=9.0) + 0.10

    vrule(s, SW / 2 - 0.01, top, 3.85, color=HAIR)
    for side, x in ((d["left"], M), (d["right"], 7.15)):
        w = 5.05
        tf = textbox(s, x, top, 1.2, 0.8)
        para(tf, side["cjk"], 38, color=RGBColor(0xC6, 0xC0, 0xB2), first=True)
        tf = textbox(s, x + 1.05, top + 0.10, w - 1.05, 0.5)
        para(tf, side["name"], 23, bold=True, color=GREEN, spc=1.2, first=True)
        tf = textbox(s, x + 1.05, top + 0.66, w - 1.05, 0.3)
        para(tf, side["lead"], 11, color=MUTED, spc=2.4, first=True)

        y = bullets(s, side["points"], x, top + 1.28, w, size=15.5, gap=0.24,
                    indent=0.38)
        rule(s, x, y + 0.16, 0.8, color=HAIR)
        ah = text_h(side["accent"], 14.5, w, 1.32, italic=True)
        tf = textbox(s, x, y + 0.38, w, ah + 0.1)
        para(tf, side["accent"], 14.5, color=INK_SOFT, italic=True, line=1.32,
             first=True)
    notes(s, d["notes"])
    footer(s, num, d["kicker"])
    return s


def s_two_terms(prs, d, num):
    s = blank(prs)
    background(s, "paper-clean.png")
    top = head(s, d["kicker"], d["title"], width=9.5) + 0.05

    if d.get("img"):                       # тайцзи — печатью в правом верхнем углу
        pic_fit(s, d["img"], SW - M - 1.45, 0.72, 1.45, 2.0)

    for side, x in ((d["left"], M), (d["right"], 7.15)):
        w = 5.05
        tf = textbox(s, x, top, 1.0, 0.85)
        para(tf, side["cjk"], 42, color=RGBColor(0xBF, 0xB8, 0xA8), first=True)
        tf = textbox(s, x + 1.02, top + 0.10, w - 1.02, 0.5)
        para(tf, side["name"], 26, bold=True, color=GREEN, spc=2.0, first=True)
        tf = textbox(s, x + 1.02, top + 0.68, w - 1.02, 0.3)
        para(tf, side["sub"], 11.5, color=MUTED, spc=2.2, first=True)
        rule(s, x, top + 1.10, w, color=HAIR, h=0.009)
        bullets(s, side["points"], x, top + 1.30, w, size=15, gap=0.22, indent=0.36)

    if d.get("highlight"):
        q = "«%s»" % d["highlight"]
        h = text_h(q, 15.5, SW - 2 * M, 1.34, italic=True)
        top2 = SH - 0.92 - h
        rule(s, M, top2 - 0.24, SW - 2 * M, color=HAIR, h=0.009)
        tf = textbox(s, M, top2, SW - 2 * M, h + 0.1)
        para(tf, q, 15.5, color=GREEN, italic=True, align=PP_ALIGN.CENTER,
             first=True)
    notes(s, d["notes"])
    footer(s, num, d["kicker"])
    return s


def s_quote(prs, d, num):
    s = blank(prs)
    background(s, "paper-quote.png")

    tf = textbox(s, M, 1.05, 9.0, 0.3)
    para(tf, d["kicker"], 10, color=MUTED, spc=3.0, first=True)

    lines = d["quote"].split("\n")
    qh = len(lines) * 26 * 1.52 / 72.0
    qy = 2.55
    tf = textbox(s, 1.55, qy, 10.2, qh + 0.2)
    for i, ln in enumerate(lines):
        para(tf, ("«" if i == 0 else "") + ln + ("»" if i == len(lines) - 1 else ""),
             26, color=INK, line=1.52, align=PP_ALIGN.CENTER, first=(i == 0))

    y = qy + qh + 0.36
    rule(s, SW / 2 - 0.4, y, 0.8, color=BRONZE)
    tf = textbox(s, M, y + 0.24, SW - 2 * M, 0.4)
    para(tf, d["author"], 13, color=GREEN, spc=3.4, align=PP_ALIGN.CENTER, first=True)

    y += 0.92
    if d.get("extra"):
        h = text_h(d["extra"], 15, 10.2, 1.34, italic=True)
        tf = textbox(s, 1.55, y, 10.2, h + 0.1)
        para(tf, d["extra"], 15, color=INK_SOFT, italic=True,
             align=PP_ALIGN.CENTER, line=1.34, first=True)
        y += h + 0.26
    if d.get("footnote"):
        h = text_h(d["footnote"], 12.5, 10.2, 1.3)
        tf = textbox(s, 1.55, y, 10.2, h + 0.1)
        para(tf, d["footnote"], 12.5, color=MUTED, align=PP_ALIGN.CENTER,
             line=1.3, first=True)
    notes(s, d["notes"])
    footer(s, num, d["kicker"])
    return s


def s_timeline(prs, d, num):
    s = blank(prs)
    background(s, "paper-clean.png")
    top = head(s, d["kicker"], d["title"], width=9.5) + 0.14

    cw = (SW - 2 * M) / 3
    for i, (age, txt) in enumerate(d["steps"]):
        col, row = i % 3, i // 3
        x = M + col * cw
        y = top + row * 1.98
        tf = textbox(s, x, y, cw - 0.5, 0.75)
        para(tf, age, 40, bold=True, color=GREEN, spc=0.6, first=True)
        tf = textbox(s, x, y + 0.70, 0.9, 0.3)
        para(tf, "лет", 11.5, color=BRONZE, spc=2.0, first=True)
        rule(s, x, y + 1.08, cw - 0.6, color=HAIR, h=0.009)
        tf = textbox(s, x, y + 1.24, cw - 0.55, 0.8)
        para(tf, txt, 14.5, color=INK, line=1.32, first=True)

    if d.get("footnote"):
        tf = textbox(s, M, SH - 1.06, SW - 2 * M, 0.5)
        para(tf, d["footnote"], 13, color=MUTED, italic=True, first=True)
    notes(s, d["notes"])
    footer(s, num, d["kicker"])
    return s


def s_ladder(prs, d, num):
    s = blank(prs)
    background(s, "paper.png")
    y = head(s, d["kicker"], d["title"], width=9.5) + 0.16

    for i, (nm, name, desc) in enumerate(d["steps"]):
        x = M + i * 0.40
        tf = textbox(s, x, y + 0.02, 0.8, 0.45)
        para(tf, nm, 19, color=BRONZE, spc=1.6, first=True)
        tf = textbox(s, x + 0.95, y, 3.3, 0.45)
        para(tf, name, 19, bold=True, color=GREEN, spc=1.2, first=True)
        dx = M + 5.35                      # колонка пояснений выровнена
        dw = SW - M - dx
        dh = text_h(desc, 15, dw, 1.3)
        tf = textbox(s, dx, y + 0.03, dw, dh + 0.1)
        para(tf, desc, 15, color=INK, line=1.3, first=True)
        rule(s, x, y + max(0.52, dh + 0.14), SW - M - x, color=HAIR, h=0.008)
        y += max(0.52, dh + 0.14) + 0.44
    notes(s, d["notes"])
    footer(s, num, d["kicker"])
    return s


def s_five(prs, d, num):
    s = blank(prs)
    background(s, "paper-clean.png")
    y = head(s, d["kicker"], d["title"], y=0.66, width=9.5)
    y = lead_line(s, d["lead"], M, y - 0.06, 9.5, size=15.5) + 0.02

    step = 0.83
    for cjk, name, sub, desc in d["cards"]:
        tf = textbox(s, M, y - 0.14, 0.85, 0.7)
        para(tf, cjk, 29, color=RGBColor(0xC2, 0xBB, 0xAC), first=True)
        tf = textbox(s, M + 0.92, y - 0.06, 2.6, 0.45)
        para(tf, name, 20, bold=True, color=GREEN, spc=1.4, first=True)
        tf = textbox(s, M + 0.92, y + 0.32, 2.9, 0.3)
        para(tf, sub, 11.5, color=MUTED, first=True)
        dx = M + 4.05
        dw = SW - M - dx
        tf = textbox(s, dx, y + 0.02, dw, 0.6)
        para(tf, desc, 15, color=INK, line=1.28, first=True)
        rule(s, M, y + step - 0.17, SW - 2 * M, color=HAIR, h=0.008)
        y += step
    notes(s, d["notes"])
    footer(s, num, d["kicker"])
    return s


def s_advice(prs, d, num):
    s = blank(prs)
    background(s, "paper.png")
    y0 = head(s, d["kicker"], d["title"], y=0.66, width=9.5)
    y0 = lead_line(s, d["lead"], M, y0 - 0.06, 9.5, size=15) + 0.04

    half = (len(d["items"]) + 1) // 2
    cw, step = 5.55, 0.84
    for i, it in enumerate(d["items"]):
        col, row = (0, i) if i < half else (1, i - half)
        x = M + col * (cw + 0.60)
        y = y0 + row * step
        tf = textbox(s, x, y + 0.04, 0.5, 0.35)
        para(tf, "%02d" % (i + 1), 12, color=BRONZE, spc=1.0, first=True)
        tf = textbox(s, x + 0.55, y, cw - 0.55, 0.7)
        para(tf, it, 14.5, color=INK, line=1.3, first=True)
        rule(s, x, y + step - 0.19, cw, color=HAIR, h=0.008)
    notes(s, d["notes"])
    footer(s, num, d["kicker"])
    return s


def s_conclusion(prs, d, num):
    s = blank(prs)
    background(s, "paper-clean.png")
    top = head(s, d["kicker"], d["title"], width=9.5) + 0.12

    for i, (n, txt) in enumerate(d["blocks"]):
        col, row = i % 2, i // 2
        x = M + col * 5.95
        y = top + row * 1.62
        tf = textbox(s, x, y, 0.7, 0.6)
        para(tf, n, 34, bold=True, color=RGBColor(0xC2, 0xBB, 0xAC), first=True)
        tf = textbox(s, x + 0.80, y + 0.08, 4.7, 1.3)
        para(tf, txt, 15.5, color=INK, line=1.34, first=True)

    highlight(s, d["highlight"], size=15)
    notes(s, d["notes"])
    footer(s, num, d["kicker"])
    return s


def s_end(prs, d, num):
    s = blank(prs)
    background(s, "section.png")
    rect(s, 0, 2.55, SW, 2.60, PAPER, transparency=0.08)
    rule(s, 0, 2.55, SW, color=HAIR, h=0.010)
    rule(s, 0, 5.15, SW, color=HAIR, h=0.010)

    tf = textbox(s, M, 2.92, 11.4, 1.5)
    for i, ln in enumerate(d["title"].split("\n")):
        para(tf, ln, 36, bold=True, color=INK, spc=3.0, line=1.1,
             align=PP_ALIGN.CENTER, first=(i == 0))

    rule(s, SW / 2 - 0.45, 4.52, 0.9, color=BRONZE)

    tf = textbox(s, 2.2, 5.62, 8.9, 1.0)
    para(tf, "«%s»" % d["quote"], 14.5, color=INK_SOFT, italic=True,
         align=PP_ALIGN.CENTER, line=1.45, first=True)
    para(tf, d["quote_author"], 10, color=MUTED, spc=2.6,
         align=PP_ALIGN.CENTER, before=9, caps=True)
    notes(s, d["notes"])
    return s


LAYOUTS = {
    "cover": s_cover, "contents": s_contents, "section": s_section,
    "text_image": s_text_image, "points": s_points,
    "two_portraits": s_two_portraits, "compare": s_compare,
    "two_terms": s_two_terms, "quote": s_quote, "timeline": s_timeline,
    "ladder": s_ladder, "five": s_five, "advice": s_advice,
    "conclusion": s_conclusion, "end": s_end,
}


def main():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)

    for i, d in enumerate(content.SLIDES, 1):
        LAYOUTS[d["layout"]](prs, d, i)

    prs.core_properties.title = content.DECK["title"]
    prs.core_properties.subject = content.DECK["subject"]

    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, "kitajskaya-filosofiya.pptx")
    prs.save(path)
    print("готово:", path, "— слайдов:", len(prs.slides.__iter__.__self__._sldIdLst))


if __name__ == "__main__":
    main()
