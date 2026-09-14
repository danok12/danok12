# -*- coding: utf-8 -*-
"""Сборка файла с текстом выступления из того же контента, что и презентация.

Речь берётся из поля notes каждого слайда — то есть текст в файле и заметки
докладчика в .pptx всегда совпадают. Правится в одном месте: content_*.py.

    python3 tools/make_speech.py --content content_confucius --out rech-konfucianstvo
"""
import argparse
import importlib
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
OUT = os.path.join(ROOT, "out")

INK = RGBColor(0x2A, 0x2A, 0x24)
GREEN = RGBColor(0x2C, 0x38, 0x2F)
MUTED = RGBColor(0x80, 0x7B, 0x70)

WPM = 135          # спокойный темп рассказа, слов в минуту


def on_screen(d):
    """Короткая сводка того, что видит зал, — чтобы не терять место в докладе."""
    L = d["layout"]
    if L == "cover":
        return "%s · %s" % (d["title"].replace("\n", " "), d["subtitle"])
    if L == "contents":
        return " · ".join(t for _, t, _ in d["items"])
    if L == "section":
        return d["title"] + " — " + d["subtitle"]
    if L in ("text_image", "points"):
        return " · ".join(d["points"])
    if L == "timeline":
        return " · ".join("%s — %s" % (a, t) for a, t in d["steps"])
    if L == "ladder":
        return " · ".join("%s: %s" % (n, t) for _, n, t in d["steps"])
    if L == "quote":
        return "цитата — «%s…»" % d["quote"].split("\n")[0]
    if L == "five":
        return " · ".join(n for _, n, _, _ in d["cards"])
    if L == "advice":
        return " · ".join(d["items"])
    if L == "conclusion":
        return " · ".join(t for _, t in d["blocks"])
    if L == "end":
        return d["title"].replace("\n", " ")
    return ""


def heading_of(d):
    if d.get("title"):
        return d["title"].replace("\n", " ")
    if d["layout"] == "quote":
        return "Цитата"
    return d.get("kicker", "")


def style_run(run, size, *, bold=False, italic=False, color=INK, caps=False):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    if caps:
        run.font.all_caps = True


def para(doc, text, size=12, *, bold=False, italic=False, color=INK,
         align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=6, line=1.4, caps=False):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    style_run(p.add_run(text), size, bold=bold, italic=italic, color=color, caps=caps)
    return p


def build(module="content_confucius", name="rech"):
    src = importlib.import_module(module)
    slides = src.SLIDES

    words = sum(len(s.get("notes", "").split()) for s in slides)
    minutes = max(1, round(words / WPM))

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.0)

    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(12)

    para(doc, src.DECK["title"].upper(), 22, bold=True, after=2)
    para(doc, "Текст выступления", 13, color=MUTED, after=2)
    para(doc, "%d слайдов · примерно %d минут · около %d слов"
         % (len(slides), minutes, words), 10.5, color=MUTED, after=18)

    for i, d in enumerate(slides, 1):
        para(doc, "СЛАЙД %d · %s" % (i, heading_of(d)), 13, bold=True,
             color=GREEN, before=16, after=2)
        scr = on_screen(d)
        if scr:
            para(doc, "На экране: " + scr, 9.5, italic=True, color=MUTED, after=8)
        for block in (d.get("notes") or "").split("\n\n"):
            block = block.strip()
            if block:
                para(doc, block, 12, after=8)

    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, name + ".docx")
    doc.save(path)
    print("готово:", path, "— слайдов: %d, слов: %d, ~%d мин" % (len(slides), words, minutes))
    return path


def main():
    ap = argparse.ArgumentParser(description="Сборка текста выступления")
    ap.add_argument("--content", default="content_confucius")
    ap.add_argument("--out", default="rech")
    a = ap.parse_args()
    build(a.content, a.out)


if __name__ == "__main__":
    main()
