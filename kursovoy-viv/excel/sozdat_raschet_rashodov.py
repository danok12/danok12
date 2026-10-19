# -*- coding: utf-8 -*-
"""Книга Excel «Расчёт расходов воды (СП 30)» для контрольной: вписать данные задания — считает всё сама.

Повторяет таблицу преподавателя (Глушков Я.А., занятие «Расчет расходов воды», НИУ МГСУ 2026):
U и N по квартирам или по заселённости, нормы по табл. А.2, суточные, среднечасовые, секундные и часовые
расходы для tot / h / c, α по табл. Б.2 линейной интерполяцией y = y1 + (x − x1)(y2 − y1)/(x2 − x1)
(или ближайшее меньшее, как в примерах методички 2024), при P > 0,1 и N ≤ 200 — по табл. Б.1.
Таблицы Б.1 и Б.2 — из calc/sp30_pril_B.json (извлечены из СП, сверены с примерами методички).

Запуск: python sozdat_raschet_rashodov.py  ->  «Расчёт расходов воды (СП 30).xlsx» рядом со скриптом.
"""
import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
B = json.load(open(os.path.join(HERE, "..", "calc", "sp30_pril_B.json"), encoding="utf-8"))
B2 = sorted((float(k), v) for k, v in B["b2"].items())
B1_N = sorted(int(k) for k in B["b1"])
B1_P = B["b1_P"]
ap = argparse.ArgumentParser(description="Книга «Расчёт расходов воды (СП 30)»; без аргументов — пример с занятия")
ap.add_argument("--U", type=float, help="число жителей, если дано в задании готовым числом")
ap.add_argument("--N", type=int, help="число приборов (с холодной водой), если дано в задании")
ap.add_argument("--Nh", type=int, help="число приборов с горячей водой (если не задано — равно N)")
ap.add_argument("--out", help="имя файла (рядом со скриптом)")
ARGS = ap.parse_args()
OUT = os.path.join(HERE, ARGS.out or "Расчёт расходов воды (СП 30).xlsx")

# табл. А.2 СП 30.13330.2020, п. 1 «Жилые здания квартирного типа», нормы на 1 жителя (calc/normy.py)
A2 = [  # наименование, q_u tot, q_u h, q_hr,u tot, q_hr,u h, q0 tot, q0 (c, h), q0,hr tot, q0,hr (c, h), T
    ("с водопроводом и канализацией без ванн", 70, 0, 5.0, 0, 0.2, 0.2, 50, 50, 24),
    ("с ваннами и водонагревателями на твёрдом топливе", 110, 0, 8.1, 0, 0.3, 0.3, 300, 300, 24),
    ("с ваннами и газовыми водонагревателями", 120, 0, 8.7, 0, 0.3, 0.3, 300, 300, 24),
    ("с ЦГВ, с умывальниками, мойками и душами", 130, 50, 8.2, 4.5, 0.2, 0.14, 100, 60, 24),
    ("с сидячими ваннами, оборудованными душами", 160, 65, 10.3, 5.7, 0.3, 0.2, 300, 200, 24),
    ("с ваннами длиной 1500–1700 мм, оборудованными душами", 180, 70, 11.6, 6.5, 0.3, 0.2, 300, 200, 24),
    ("свои значения (вписать в эту строку)", None, None, None, None, None, None, None, None, None),
]

SH = "Расчёт"
F_NAME = "Times New Roman"
f_base = Font(name=F_NAME, size=12)
f_bold = Font(name=F_NAME, size=12, bold=True)
f_head = Font(name=F_NAME, size=13, bold=True)
f_title = Font(name=F_NAME, size=15, bold=True)
f_norm = Font(name=F_NAME, size=12, color="0000FF")         # нормы из таблиц — можно исправить
f_note = Font(name=F_NAME, size=10, italic=True, color="595959")
f_res = Font(name=F_NAME, size=12, bold=True)
fill_in = PatternFill("solid", fgColor="FFF2A8")           # вписать данные
fill_norm = PatternFill("solid", fgColor="DDEBF7")         # нормы А.2 (как у преподавателя — голубые)
fill_alpha = PatternFill("solid", fgColor="C6EFCE")        # α по Б.2 (как у преподавателя — зелёные)
fill_res = PatternFill("solid", fgColor="E2EFDA")
fill_head = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="808080")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)

wb = Workbook()
ws = wb.active
ws.title = SH


def put(ref, val, font=f_base, fill=None, fmt=None, al=None, border=True):
    c = ws[ref]
    c.value = val
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    c.alignment = al or (left if isinstance(val, str) and not str(val).startswith("=") else center)
    if border:
        c.border = box
    return c


def header(ref, text):
    put(ref, text, f_head, border=False, al=Alignment(horizontal="left", vertical="center"))


# ---------------- оформление листа ----------------
for col, w in zip("ABCDEFGHIJKLMNOPQRS", (2, 58, 13, 13, 13, 13, 13, 13, 13, 13, 12, 11, 11, 11, 11, 11, 11, 11, 12)):
    ws.column_dimensions[col].width = w
ws.sheet_view.zoomScale = 90

put("B1", "Расчёт расходов воды по СП 30.13330.2020 (п. 5, табл. А.2, Б.1, Б.2)", f_title, border=False)
ws.merge_cells("B1:S1")
put("B2", "Вписывайте данные только в жёлтые ячейки. Голубые — нормы из табл. А.2, подставляются сами по виду "
          "здания (можно исправить вручную). Зелёные — α по табл. Б.2/Б.1 и результаты. Ниже, в разделе 7, — "
          "запись решения с подставленными числами.", f_note, border=False)
ws.merge_cells("B2:S2")
ws.row_dimensions[2].height = 30

# ---------------- 1. исходные данные ----------------
header("B4", "1. Исходные данные")
put("B5", "Число этажей, n_эт")
put("C5", 12, f_base, fill_in, "0")
put("B6", "Число секций (квартиры ниже — на этаже одной секции)")
put("C6", 1, f_base, fill_in, "0")
put("B7", "Заселённость, чел./кв. — если дана в задании (иначе оставить пустой: U считается по типам квартир)")
put("C7", 3.4, f_base, fill_in, "0.0##")
put("B8", "Вид жилого здания (табл. А.2 СП 30.13330.2020) — выбрать из списка")
ws.merge_cells("C8:I8")
put("C8", A2[5][0], f_base, fill_in)
for col in "DEFGHI":
    ws[col + "8"].fill = fill_in
    ws[col + "8"].border = box
put("B9", "Как брать α по табл. Б.2 — выбрать из списка")
put("C9", "интерполяция", f_base, fill_in)
ws.merge_cells("C9:D9")
put("E9", "интерполяция — как в таблице преподавателя; «ближайшее меньшее» — как в примерах методички 2024",
    f_note, border=False)
ws.merge_cells("E9:S9")

# ---------------- 2. квартиры ----------------
header("B11", "2. Квартиры на этаже (в одной секции)")
heads = ["Тип квартиры", "Квартир на этаже, n_кв", "Комнат, K", "Жителей в квартире, N_кв.жит (= K + 1)",
         "Приборов в квартире всего, n_пр (ХВС)", "из них с горячей водой (ГВС)", "Жителей на этаже",
         "Приборов ХВС на этаже", "Приборов ГВС на этаже"]
for col, h in zip("BCDEFGHIJ", heads):
    put(col + "12", h, f_bold, fill_head, al=center)
ws.row_dimensions[12].height = 64
TIPY = [("1-комнатная", None, 1), ("2-комнатная", 4, 2), ("3-комнатная", 4, 3), ("4-комнатная", None, 4),
        ("другая", None, None)]
for k, (nm, n_kv, komn) in enumerate(TIPY):
    r = 13 + k
    put(f"B{r}", nm, f_base, fill_in)
    put(f"C{r}", n_kv, f_base, fill_in, "0;-0;\"\"")
    put(f"D{r}", komn, f_base, fill_in, "0;-0;\"\"")
    put(f"E{r}", f'=IF(OR(C{r}="",D{r}=""),0,D{r}+1)', f_base, fill_in, "0.0##;-0;\"\"")
    put(f"F{r}", 4 if n_kv else None, f_base, fill_in, "0;-0;\"\"")
    put(f"G{r}", 3 if n_kv else None, f_base, fill_in, "0;-0;\"\"")
    put(f"H{r}", f"=C{r}*E{r}", fmt="0.0##;-0;\"\"")
    put(f"I{r}", f"=C{r}*F{r}", fmt="0;-0;\"\"")
    put(f"J{r}", f"=C{r}*G{r}", fmt="0;-0;\"\"")
ws["E13"].comment = Comment("По умолчанию жителей = комнат + 1 (как у преподавателя: 2-к — 3 чел., 3-к — 4 чел.). "
                            "Если в задании другое — впишите число поверх формулы.", "расчёт")
ws["F13"].comment = Comment("Квартира с ванной: мойка, умывальник, ванна, унитаз — 4 прибора с холодной водой, "
                            "из них 3 с горячей (унитаз — только холодная).", "расчёт")
put("B18", "Итого на этаже секции", f_bold)
for col in "CHIJ":
    put(f"{col}18", f"=SUM({col}13:{col}17)", f_bold, fmt="0.0##;-0;\"\"" if col == "H" else "0")

# ---------------- 3. U и N ----------------
header("B20", "3. Число водопотребителей U и приборов N")
put("B21", "U по типам квартир = Σ(N_кв.жит · n_кв) · n_эт · n_сек, чел.")
put("C21", "=H18*C5*C6", fmt="0.0##")
put("B22", "U по заселённости = заселённость · n_кв · n_эт · n_сек, чел.")
put("C22", '=IF(C7="","",C7*C18*C5*C6)', fmt="0.0##")
put("B23", "U, принятое в расчёте (задано; иначе по заселённости — округление вверх), чел.", f_res)
put("C23", '=IF(D23<>"",D23,IF(C7="",C21,ROUNDUP(C22,0)))', f_res, fill_res, "0")
put("B24", "N — приборов с холодной водой (ХВС) = Σ(n_пр · n_кв) · n_эт · n_сек, шт.", f_res)
put("C24", '=IF(D24<>"",D24,I18*C5*C6)', f_res, fill_res, "0")
put("B25", "N_h — приборов с горячей водой (ГВС), шт.", f_res)
put("C25", '=IF(D25<>"",D25,IF(D24<>"",D24,J18*C5*C6))', f_res, fill_res, "0")
put("D20", "Задано в задании", f_bold, fill_head, al=center)
for r_, v_ in ((23, ARGS.U), (24, ARGS.N), (25, ARGS.Nh)):
    put(f"D{r_}", v_, f_base, fill_in, "0;-0;\"\"")
put("E23", "← если U, N, N_h даны в задании готовыми числами — вписать сюда (разделы 1–2 тогда не нужны); "
           "N_h пусто — принимается равным N", f_note, border=False)
ws.merge_cells("E23:S23")

# ---------------- 4. нормы А.2 ----------------
header("B27", "4. Нормы водопотребления по табл. А.2 СП 30.13330.2020 (на 1 жителя / на 1 прибор)")
for col, h in zip("BCDEF", ["Норма", "tot (общий)", "h (горячей)", "c (холодной)", "ед. изм."]):
    put(col + "28", h, f_bold, fill_head, al=center)
NORMY = [  # строка, подпись, столбец tot, столбец h в листе А.2, как получить c, ед.
    (29, "q_u — расход в сутки наибольшего водопотребления", "B", "C", "minus", "л/сут"),
    (30, "q_hr,u — расход в час наибольшего водопотребления", "D", "E", "minus", "л/ч"),
    (31, "q_0 — секундный расход воды прибором", "F", "G", "same", "л/с"),
    (32, "q_0,hr — часовой расход воды прибором", "H", "I", "same", "л/ч"),
    (33, "T — время водопотребления", "J", "J", "same", "ч"),
]
for r, txt, ctot, ch, how, ed in NORMY:
    put(f"B{r}", txt)
    put(f"C{r}", f"=INDEX(A2_{ctot},MATCH($C$8,A2_IMYA,0))", f_norm, fill_norm, "0.0##")
    put(f"D{r}", f"=INDEX(A2_{ch},MATCH($C$8,A2_IMYA,0))", f_norm, fill_norm, "0.0##")
    put(f"E{r}", f"=C{r}-D{r}" if how == "minus" else f"=D{r}", fmt="0.0##")
    put(f"F{r}", ed, al=center)
ws["E31"].comment = Comment("В табл. А.2 q_0 и q_0,hr холодной и горячей воды — одна графа («холодной или горячей»).",
                            "расчёт")
put("B34", "q_u^c = q_u^tot − q_u^h; q_hr,u^c = q_hr,u^tot − q_hr,u^h; q_0^c = q_0^h; q_0,hr^c = q_0,hr^h "
           "(табл. А.2). Нормы можно вписать вручную поверх голубых ячеек.", f_note, border=False)
ws.merge_cells("B34:S34")

# ---------------- 5. расчёт ----------------
header("B36", "5. Расчёт расходов воды (п. 5 СП 30.13330.2020)")
COLS5 = ["", "T, ч", "U, чел.", "N, шт.", "q_u, л/сут", "q_hr,u, л/ч", "q_0, л/с", "q_0,hr, л/ч",
         "q_u,m — сут., м³/сут", "q_T — ср.-час., м³/ч", "P", "NP", "α", "q — сек., л/с",
         "P_hr", "NP_hr", "α_hr", "q_hr — час., м³/ч"]
for col, h in zip("BCDEFGHIJKLMNOPQRS", COLS5):
    put(col + "37", h, f_bold, fill_head, al=center)
ws.row_dimensions[37].height = 46
STROKI = [(38, "tot (общий)", "C", "$C$24", 1), (39, "h (горячей)", "D", "$C$25", 2), (40, "c (холодной)", "E", "$C$24", 3)]
for r, nm, nc, n_ref, k_alpha in STROKI:
    put(f"B{r}", nm, f_bold)
    put(f"C{r}", f"={nc}33", fmt="0")
    put(f"D{r}", "=$C$23", fmt="0")
    put(f"E{r}", f"={n_ref}", fmt="0")
    put(f"F{r}", f"={nc}29", fmt="0.0##")
    put(f"G{r}", f"={nc}30", fmt="0.0##")
    put(f"H{r}", f"={nc}31", fmt="0.0##")
    put(f"I{r}", f"={nc}32", fmt="0.0##")
    put(f"J{r}", f"=F{r}*D{r}/1000", f_res, fill_res, "0.000")
    put(f"K{r}", f"=IF(C{r}=0,0,J{r}/C{r})", f_res, fill_res, "0.000")
    put(f"L{r}", f"=IF(OR(E{r}=0,H{r}=0),0,G{r}*D{r}/(H{r}*E{r}*3600))", fmt="0.00000")
    put(f"M{r}", f"=E{r}*L{r}", fmt="0.000")
    put(f"N{r}", f"='α'!AA{1 + k_alpha}", f_base, fill_alpha, "0.000")
    put(f"O{r}", f"=IF(L{r}=0,0,5*H{r}*N{r})", f_res, fill_res, "0.000")
    put(f"P{r}", f"=IF(I{r}=0,0,3600*L{r}*H{r}/I{r})", fmt="0.00000")
    put(f"Q{r}", f"=E{r}*P{r}", fmt="0.000")
    put(f"R{r}", f"='α'!AA{4 + k_alpha}", f_base, fill_alpha, "0.000")
    put(f"S{r}", f"=IF(P{r}=0,0,0.005*I{r}*R{r})", f_res, fill_res, "0.000")
put("B41", "P = q_hr,u·U/(q_0·N·3600); q = 5·q_0·α; P_hr = 3600·P·q_0/q_0,hr; q_hr = 0,005·q_0,hr·α_hr; "
           "q_u,m = q_u·U/1000; q_T = q_u,m/T. Для c: q_hr,u^c = q_hr,u^tot − q_hr,u^h.", f_note, border=False)
ws.merge_cells("B41:S41")

# ---------------- 6. α ----------------
header("B43", "6. Определение α: y = y₁ + (x − x₁)·(y₂ − y₁)/(x₂ − x₁)  (при P > 0,1 и N ≤ 200 — по табл. Б.1)")
for col, h in zip("BCDEFGHIJ", ["Случай", "Таблица", "NP₁ (x₁)", "α₁ (y₁)", "NP₂ (x₂)", "α₂ (y₂)", "NP (x)",
                                 "α (y)", "Примечание"]):
    put(col + "44", h, f_bold, fill_head, al=center)
SLUCHAI = ["секундный tot", "секундный h", "секундный c", "часовой tot", "часовой h", "часовой c"]
for k, nm in enumerate(SLUCHAI):
    r, hr = 45 + k, k + 2          # строка на листе α
    put(f"B{r}", nm)
    put(f"C{r}", f"='α'!AB{hr}", al=center)
    put(f"D{r}", f"='α'!G{hr}", fmt="0.000")
    put(f"E{r}", f"='α'!H{hr}", fmt="0.000")
    put(f"F{r}", f"='α'!I{hr}", fmt="0.000")
    put(f"G{r}", f"='α'!J{hr}", fmt="0.000")
    put(f"H{r}", f"='α'!D{hr}", fmt="0.000")
    put(f"I{r}", f"='α'!AA{hr}", f_res, fill_alpha, "0.000")
    put(f"J{r}", f"='α'!AC{hr}", f_note)
    ws.merge_cells(f"J{r}:S{r}")

# ---------------- 7. запись решения ----------------
header("B52", "7. Запись решения (с подставленными числами — переписать в работу)")


def fx(ref, n):
    return f"FIXED({ref},{n},TRUE)"


ZAP = [
    f'=IF($D$23<>"","U = "&C23&" чел. (задано)",IF($C$7="","U = Σ(N_кв.жит·n_кв)·n_эт·n_сек = "&{fx("H18", 1)}&"·"&C5&"·"&C6&" = "&{fx("C23", 0)}&" чел.",'
    f'"U = заселённость·n_кв·n_эт·n_сек = "&{fx("C7", 2)}&"·"&C18&"·"&C5&"·"&C6&" = "&{fx("C22", 2)}&" ≈ "&C23&" чел."))',
    f'=IF($D$24<>"","N = "&C24&" шт. (задано);   N_h = "&C25&" шт."&IF($D$25="",IF(C25=C24," (принято равным N)",""),'
    f'" (задано)"),"N = Σ(n_пр·n_кв)·n_эт·n_сек = "&I18&"·"&C5&"·"&C6&" = "&C24&" шт.;   N_h = "&J18&"·"&C5&"·"&C6&" = "&C25&" шт.")',
    f'="Суточные: q_u,m^tot = q_u^tot·U/1000 = "&{fx("F38", 1)}&"·"&D38&"/1000 = "&{fx("J38", 3)}&" м³/сут;  '
    f'q_u,m^c = (q_u^tot − q_u^h)·U/1000 = ("&{fx("C29", 1)}&" − "&{fx("D29", 1)}&")·"&D40&"/1000 = "&{fx("J40", 3)}&" м³/сут;  '
    f'q_u,m^h = "&{fx("F39", 1)}&"·"&D39&"/1000 = "&{fx("J39", 3)}&" м³/сут"',
    f'="Среднечасовые: q_T^tot = q_u^tot·U/(1000·T) = "&{fx("J38", 3)}&"/"&C38&" = "&{fx("K38", 3)}&" м³/ч;  '
    f'q_T^c = "&{fx("J40", 3)}&"/"&C40&" = "&{fx("K40", 3)}&" м³/ч;  q_T^h = "&{fx("J39", 3)}&"/"&C39&" = "&{fx("K39", 3)}&" м³/ч"',
]
for r_, nm, sym in ((38, "tot", "tot"), (40, "c", "c"), (39, "h", "h")):
    k_ = {"tot": 45, "h": 46, "c": 47}[nm]
    q_hr_txt = (f'"("&{fx("C30", 2)}&" − "&{fx("D30", 2)}&")"' if nm == "c" else f'{fx(f"G{r_}", 2)}')
    ZAP.append(
        f'="P^{sym} = q_hr,u^{sym}·U/(q_0^{sym}·N·3600) = "&{q_hr_txt}&"·"&D{r_}&"/("&{fx(f"H{r_}", 2)}&"·"&E{r_}&"·3600) = "'
        f'&{fx(f"L{r_}", 5)}&";  NP^{sym} = "&E{r_}&"·"&{fx(f"L{r_}", 5)}&" = "&{fx(f"M{r_}", 3)}&" → α^{sym} = "'
        f'&{fx(f"N{r_}", 3)}&" (табл. "&C{k_}&");  q^{sym} = 5·q_0^{sym}·α^{sym} = 5·"&{fx(f"H{r_}", 2)}&"·"&{fx(f"N{r_}", 3)}'
        f'&" = "&{fx(f"O{r_}", 3)}&" л/с"')
for r_, nm, sym in ((38, "tot", "tot"), (40, "c", "c"), (39, "h", "h")):
    k_ = {"tot": 48, "h": 49, "c": 50}[nm]
    ZAP.append(
        f'="P_hr^{sym} = 3600·P^{sym}·q_0^{sym}/q_0,hr^{sym} = 3600·"&{fx(f"L{r_}", 5)}&"·"&{fx(f"H{r_}", 2)}&"/"&{fx(f"I{r_}", 0)}'
        f'&" = "&{fx(f"P{r_}", 5)}&";  NP_hr^{sym} = "&E{r_}&"·"&{fx(f"P{r_}", 5)}&" = "&{fx(f"Q{r_}", 3)}&" → α_hr^{sym} = "'
        f'&{fx(f"R{r_}", 3)}&" (табл. "&C{k_}&");  q_hr^{sym} = 0,005·q_0,hr^{sym}·α_hr^{sym} = 0,005·"&{fx(f"I{r_}", 0)}'
        f'&"·"&{fx(f"R{r_}", 3)}&" = "&{fx(f"S{r_}", 3)}&" м³/ч"')
for k, f in enumerate(ZAP):
    r = 53 + k
    put(f"B{r}", f, f_base, border=False, al=Alignment(horizontal="left", vertical="center", wrap_text=True))
    ws.merge_cells(f"B{r}:S{r}")
    ws.row_dimensions[r].height = 34
put(f"B{53 + len(ZAP) + 1}", "Пример заполнения — данные с занятия (12 этажей, по 4 двух- и трёхкомнатных квартиры, "
    "заселённость 3,4 чел./кв., дом с ваннами 1500–1700 мм): U = 327 чел., N = 384, N_h = 288. Для контрольной — "
    "заменить жёлтые ячейки своими данными.", f_note, border=False)
ws.merge_cells(f"B{53 + len(ZAP) + 1}:S{53 + len(ZAP) + 1}")

# ---------------- списки выбора ----------------
dv_vid = DataValidation(type="list", formula1="=A2_IMYA", allow_blank=False)
dv_alpha = DataValidation(type="list", formula1='"интерполяция,ближайшее меньшее"', allow_blank=False)
ws.add_data_validation(dv_vid)
ws.add_data_validation(dv_alpha)
dv_vid.add("C8")
dv_alpha.add("C9")
ws.freeze_panes = "C4"

# ---------------- лист α (вспомогательный) ----------------
wa = wb.create_sheet("α")
ha = ["Случай", "N", "P", "NP", "по Б.1?", "i (Б.2)", "NP1", "α1", "NP2", "α2", "α по Б.2", "r (Б.1)", "r2", "c (Б.1)",
      "c2", "N1", "N2", "P1", "P2", "α11", "α12", "α21", "α22", "tP", "tN", "α по Б.1", "α", "табл.", "примечание"]
for j, h in enumerate(ha):
    c = wa.cell(row=1, column=1 + j, value=h)
    c.font = f_bold
    c.fill = fill_head
SRC = [  # строка листа «Расчёт», столбцы N и P
    (38, "E", "L"), (39, "E", "L"), (40, "E", "L"), (38, "E", "P"), (39, "E", "P"), (40, "E", "P")]
MODE = f"'{SH}'!$C$9"
for k, (rr, cn, cp) in enumerate(SRC):
    r = 2 + k
    vals = {
        "A": SLUCHAI[k],
        "B": f"='{SH}'!{cn}{rr}",
        "C": f"='{SH}'!{cp}{rr}",
        "D": f"=B{r}*C{r}",
        "E": f"=AND(C{r}>0.1,B{r}<=200,B{r}>=2)",
        "F": f"=IF(D{r}<0.015,0,MATCH(D{r},B2_NP,1))",
        "G": f"=IF(F{r}=0,0,INDEX(B2_NP,F{r}))",
        "H": f"=IF(F{r}=0,0.2,INDEX(B2_A,F{r}))",
        "I": f"=IF(F{r}=0,0.015,INDEX(B2_NP,MIN(F{r}+1,{len(B2)})))",
        "J": f"=IF(F{r}=0,0.2,INDEX(B2_A,MIN(F{r}+1,{len(B2)})))",
        "K": f'=IF(OR(F{r}=0,I{r}=G{r},{MODE}="ближайшее меньшее"),H{r},H{r}+(D{r}-G{r})*(J{r}-H{r})/(I{r}-G{r}))',
        "L": f"=IF(E{r},MATCH(B{r},B1_N,1),1)",
        "M": f"=MIN(L{r}+1,{len(B1_N)})",
        "N": f"=IF(E{r},MATCH(C{r},B1_P,1),1)",
        "O": f"=MIN(N{r}+1,{len(B1_P)})",
        "P": f"=INDEX(B1_N,L{r})",
        "Q": f"=INDEX(B1_N,M{r})",
        "R": f"=INDEX(B1_P,N{r})",
        "S": f"=INDEX(B1_P,O{r})",
        "T": f"=INDEX(B1_T,L{r},N{r})",
        "U": f"=INDEX(B1_T,L{r},O{r})",
        "V": f"=INDEX(B1_T,M{r},N{r})",
        "W": f"=INDEX(B1_T,M{r},O{r})",
        "X": f"=IF(OR(S{r}=R{r},{MODE}=\"ближайшее меньшее\"),0,(MIN(C{r},S{r})-R{r})/(S{r}-R{r}))",
        "Y": f"=IF(OR(Q{r}=P{r},{MODE}=\"ближайшее меньшее\"),0,(B{r}-P{r})/(Q{r}-P{r}))",
        "Z": f"=(T{r}+(U{r}-T{r})*X{r})*(1-Y{r})+(V{r}+(W{r}-V{r})*X{r})*Y{r}",
        "AA": f"=IF(E{r},Z{r},K{r})",
        "AB": f'=IF(E{r},"Б.1","Б.2")',
        "AC": f'=IF(E{r},"P > 0,1 и N ≤ 200 — α по табл. Б.1: N = "&P{r}&"…"&Q{r}&", P = "&R{r}&"…"&S{r},'
              f'IF(F{r}=0,"NP < 0,015 — α = 0,2 (табл. Б.2)",""))',
    }
    for col, v in vals.items():
        wa[f"{col}{r}"] = v
        wa[f"{col}{r}"].font = f_base
wa.column_dimensions["A"].width = 16
for j in range(2, len(ha) + 1):
    wa.column_dimensions[wa.cell(row=1, column=j).column_letter].width = 11
wa.column_dimensions["AC"].width = 60
wa["A10"] = "Вспомогательный лист: подбор α по табл. Б.2 (интерполяция или ближайшее меньшее) и Б.1 (двойная интерполяция по N и P)."
wa["A10"].font = f_note

# ---------------- справочные таблицы ----------------
w2 = wb.create_sheet("Табл. Б.2")
w2["A1"] = "Табл. Б.2 СП 30.13330.2020: α при P ≤ 0,1 и любом N, а также при P > 0,1 и N > 200 (NP < 0,015 → α = 0,2)"
w2["A1"].font = f_bold
w2["A2"], w2["B2"] = "NP", "α"
for c in (w2["A2"], w2["B2"]):
    c.font = f_bold
    c.fill = fill_head
for k, (np_, a) in enumerate(B2):
    w2.cell(row=3 + k, column=1, value=np_).font = f_base
    w2.cell(row=3 + k, column=2, value=a).font = f_base
n2 = 2 + len(B2)

w1 = wb.create_sheet("Табл. Б.1")
w1["A1"] = "Табл. Б.1 СП 30.13330.2020: α при P > 0,1 и N ≤ 200 (строки — N, столбцы — P)"
w1["A1"].font = f_bold
w1["A2"] = "N \\ P"
w1["A2"].font = f_bold
for j, p in enumerate(B1_P):
    c = w1.cell(row=2, column=2 + j, value=p)
    c.font = f_bold
    c.fill = fill_head
for i, n_ in enumerate(B1_N):
    w1.cell(row=3 + i, column=1, value=n_).font = f_bold
    for j, a in enumerate(B["b1"][str(n_)]):
        w1.cell(row=3 + i, column=2 + j, value=a).font = f_base
n1 = 2 + len(B1_N)
last_col = chr(ord("A") + len(B1_P))

wA = wb.create_sheet("Табл. А.2")
wA["A1"] = "Табл. А.2 СП 30.13330.2020, п. 1 «Жилые здания квартирного типа» — нормы на 1 жителя"
wA["A1"].font = f_bold
heads_a = ["Жилые здания квартирного типа", "q_u tot, л/сут", "q_u h, л/сут", "q_hr,u tot, л/ч", "q_hr,u h, л/ч",
           "q_0 tot, л/с", "q_0 (c, h), л/с", "q_0,hr tot, л/ч", "q_0,hr (c, h), л/ч", "T, ч"]
for j, h in enumerate(heads_a):
    c = wA.cell(row=2, column=1 + j, value=h)
    c.font = f_bold
    c.fill = fill_head
    c.alignment = center
for i, row in enumerate(A2):
    for j, v in enumerate(row):
        c = wA.cell(row=3 + i, column=1 + j, value=v)
        c.font = f_base
        c.border = box
        if row[1] is None:
            c.fill = fill_in
wA.column_dimensions["A"].width = 58
for col in "BCDEFGHIJ":
    wA.column_dimensions[col].width = 13
nA = 2 + len(A2)
wA[f"A{nA + 2}"] = ("Источник: СП 30.13330.2020, прил. А, табл. А.2 (значения сверены с методическими указаниями "
                    "НИУ МГСУ 2024). Для другого вида здания впишите нормы в строку «свои значения» и выберите её "
                    "на листе «Расчёт».")
wA[f"A{nA + 2}"].font = f_note

# ---------------- имена диапазонов ----------------
NAMES = {
    "B2_NP": f"'Табл. Б.2'!$A$3:$A${n2}", "B2_A": f"'Табл. Б.2'!$B$3:$B${n2}",
    "B1_N": f"'Табл. Б.1'!$A$3:$A${n1}", "B1_P": f"'Табл. Б.1'!$B$2:${last_col}$2",
    "B1_T": f"'Табл. Б.1'!$B$3:${last_col}${n1}",
    "A2_IMYA": f"'Табл. А.2'!$A$3:$A${nA}",
}
for j, col in enumerate("BCDEFGHIJ"):
    NAMES[f"A2_{col}"] = f"'Табл. А.2'!${col}$3:${col}${nA}"
for nm, ref in NAMES.items():
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print(OUT)
