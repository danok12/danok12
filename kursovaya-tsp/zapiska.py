"""Пояснительная записка (Word) - технологическая карта на земляные работы, вариант 2.

Все числа берутся из модулей расчёта (raschet, kotlovan, balans, raspredelenie, mashiny),
поэтому записка всегда совпадает с книгой Excel и чертежами. Оформление - ГОСТ 7.32:
А4, поля 30/15/20/20 мм, Times New Roman 14, интервал 1,5, абзац 1,25 см; разделы - с новой
страницы, таблицы нумеруются по разделам. Знаки «—», «–», «−» не используются - дефис.
Содержание и номера страниц обновляет Word (COM), он же сохраняет PDF.

    python zapiska.py
"""
import subprocess
from pathlib import Path
from math import sqrt
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Mm, Cm
import re
import pymupdf
import raschet as R
import kotlovan as K
import balans as B
import raspredelenie as RS
import mashiny as MS

HERE = Path(__file__).resolve().parent
OUT = HERE / "Пояснительная записка. Вариант 2.docx"
IMG = HERE / "_acad" / "zapiska"
SHEETS = [  # графическая часть - листы разд. 3 методички (без первого варианта)
    ("Лист 1. План площадки с рабочими отметками и ЛНР (черновик)", 1),
    ("Лист 2. План площадки с окончательными рабочими отметками и обводами", 2),
    ("Лист 3. Продольный и поперечный разрезы площадки по котловану", 3),
    ("Лист 4. Картограмма перемещения земляных масс", 4),
    ("Лист 5. Схема производства работ по вертикальной планировке", 5),
]
FONT = "Times New Roman"
NB = " "                          # неразрывный пробел - разделитель тысяч


def num(v, n=2):
    """Число по-русски: 1 234,56; минус - дефисом."""
    s = f"{abs(float(v)):,.{n}f}".replace(",", NB).replace(".", ",")
    return ("-" if float(v) < 0 and s.strip("0,") else "") + s


def sgn(v, n=2):
    """Рабочая отметка со знаком: +0,15 / -0,25."""
    s = num(abs(v), n)
    return s if float(v) == 0 else ("+" if v > 0 else "-") + s


# ---------------- документ и стили ----------------
doc = Document()
st = doc.styles["Normal"]
st.font.name, st.font.size = FONT, Pt(14)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
pf = st.paragraph_format
pf.line_spacing, pf.space_before, pf.space_after = 1.5, Pt(0), Pt(0)
pf.first_line_indent, pf.alignment = Cm(1.25), WD_ALIGN_PARAGRAPH.JUSTIFY
for name, size in (("Heading 1", 14), ("Heading 2", 14)):
    h = doc.styles[name]
    h.font.name, h.font.size, h.font.bold, h.font.italic = FONT, Pt(size), True, False
    h.font.color.rgb = None
    rpr = h.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    hp = h.paragraph_format
    hp.first_line_indent, hp.alignment = Cm(1.25), WD_ALIGN_PARAGRAPH.LEFT
    hp.space_before, hp.space_after, hp.line_spacing = Pt(12), Pt(12), 1.5
    hp.keep_with_next = True
doc.styles["Heading 1"].paragraph_format.page_break_before = True

sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Mm(30), Mm(15), Mm(20), Mm(20)
sec.different_first_page_header_footer = True          # титульный лист - без номера


def field(par, instr):
    """Поле Word (PAGE, TOC ...) - begin / instr / separate / end."""
    for kind, text in (("begin", None), (None, instr), ("separate", None), ("end", None)):
        r = par.add_run()
        if kind:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; r._r.append(it)


fp = sec.footer.paragraphs[0]
fp.alignment, fp.paragraph_format.first_line_indent = WD_ALIGN_PARAGRAPH.CENTER, Cm(0)
field(fp, "PAGE")


UNITS = re.compile(r"(\d) (м³|м²|мм|км|кг/м³|т/м³|м|т|мин|шт\.|маш\.-см|маш\.-ч|%|дня|дней)(?=[\s,.;:)]|$)")


def nb(text):
    """Число не отрывается от единицы измерения."""
    return UNITS.sub(lambda m_: m_.group(1) + NB + m_.group(2), text)


def para(text="", bold=False, align=None, indent=True, size=None, italic=False, keep=False, space_after=0,
         space_before=0, single=False):
    p = doc.add_paragraph()
    text = nb(text)
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size: r.font.size = Pt(size)
    if align is not None: p.alignment = align
    if not indent: p.paragraph_format.first_line_indent = Cm(0)
    if keep: p.paragraph_format.keep_with_next = True
    if space_after: p.paragraph_format.space_after = Pt(space_after)
    if space_before: p.paragraph_format.space_before = Pt(space_before)
    if single: p.paragraph_format.line_spacing = 1.0
    return p


def h1(t): return doc.add_heading(t, level=1)
def h2(t): return doc.add_heading(t, level=2)


FN = [0]


def formula(text, label=True):
    """Формула по центру, номер (раздел.n) - у правого поля."""
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    ts = p.paragraph_format.tab_stops
    ts.add_tab_stop(Mm(82.5), WD_TAB_ALIGNMENT.CENTER)
    ts.add_tab_stop(Mm(165), WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.keep_with_next = True
    if label:
        FN[0] += 1
        p.add_run(f"\t{text}\t({SEC[0]}.{FN[0]})")
        return f"({SEC[0]}.{FN[0]})"
    p.add_run(f"\t{text}")


SEC, TN = [1], [0]


def section(n, title):
    SEC[0], FN[0], TN[0] = n, 0, 0
    return h1(f"{n} {title}")


def table(title, head, rows, widths, size=12, align=None, bold_last=False, head_rows=1):
    """Таблица: подпись «Таблица N.n - Название» над ней, шрифт 12, интервал 1."""
    TN[0] += 1
    cap = para(f"Таблица {SEC[0]}.{TN[0]} - {title}", indent=False, keep=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    cap.paragraph_format.space_before = Pt(6)
    allrows = head + rows
    t = doc.add_table(rows=len(allrows), cols=len(widths))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate(allrows):
        cells = t.rows[i].cells
        if isinstance(row, str):                # строка-заголовок группы или итог
            for j, w in enumerate(widths):
                cells[j].width = Mm(w)
            m_ = cells[0].merge(cells[-1])
            m_.text = ""
            p = m_.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(nb(row))
            r.font.size, r.bold = Pt(size), True
            continue
        for j, w in enumerate(widths):
            c = cells[j]
            c.width = Mm(w)
            v = row[j] if j < len(row) else ""
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            hd = i < head_rows
            a = (align[j] if align and not hd else "c")
            p.alignment = {"c": WD_ALIGN_PARAGRAPH.CENTER, "l": WD_ALIGN_PARAGRAPH.LEFT,
                           "r": WD_ALIGN_PARAGRAPH.RIGHT}[a]
            r = p.add_run(nb(str(v)))
            r.font.size = Pt(size)
            r.bold = hd or (bold_last and i == len(allrows) - 1)
            if i < head_rows + 2:
                p.paragraph_format.keep_with_next = True
        if i < head_rows:                       # шапка повторяется на каждой странице
            trp = t.rows[i]._tr.get_or_add_trPr()
            th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trp.append(th)
    para("", indent=False)
    return f"табл. {SEC[0]}.{TN[0]}"


def where(*items):
    """«где x - ...; y - ...» после формулы."""
    para("где " + ";\n".join(items).replace("\n", " ") + ".", indent=False)


# ---------------- титульный лист (прил. 1 методических указаний) ----------------
C = WD_ALIGN_PARAGRAPH.CENTER
T = dict(indent=False, single=True)
para("Министерство науки и высшего образования Российской Федерации", align=C, **T)
para("Федеральное государственное бюджетное образовательное учреждение высшего образования", align=C, **T)
para("«НАЦИОНАЛЬНЫЙ ИССЛЕДОВАТЕЛЬСКИЙ МОСКОВСКИЙ ГОСУДАРСТВЕННЫЙ СТРОИТЕЛЬНЫЙ УНИВЕРСИТЕТ»", align=C, bold=True, **T)
para("Институт строительства и архитектуры", align=C, space_before=24, **T)
para("Кафедра технологии и организации строительного производства", align=C, **T)
para("КУРСОВАЯ РАБОТА", bold=True, align=C, size=16, space_before=84, **T)
para("по дисциплине «Технологические процессы в строительстве»", align=C, space_before=6, **T)
para("Тема: «Разработка технологической карты на производство земляных работ»", align=C, bold=True,
     space_before=18, **T)
para("Вариант 2", align=C, space_before=6, **T)


def sign_block(lines, hint, before):
    for k_, ln_ in enumerate(lines):
        p = para(ln_, space_before=before if k_ == 0 else 0, **T)
        p.paragraph_format.left_indent = Mm(85)
    if hint:
        p = para(hint, size=10, italic=True, **T)
        p.paragraph_format.left_indent = Mm(85)


sign_block(["Выполнил студент", "ИГЭС, 3 курс, группа 3-8", "Буев Даниил Романович"],
           "(институт, курс, группа; Ф.И.О.)", 60)
sign_block(["Руководитель работы", "доцент Забелина О.Б."], "(должность, Ф.И.О.)", 18)
sign_block(["К защите _____________________"], "(дата, подпись руководителя)", 18)
sign_block(["Работа защищена с оценкой", "______________________________"], "", 18)
para("Москва 2026", align=C, space_before=40, **T)

# ---------------- содержание ----------------
p = para("СОДЕРЖАНИЕ", bold=True, align=C, indent=False, space_after=12)
p.paragraph_format.page_break_before = True
field(para(indent=False), 'TOC \\o "1-2" \\h \\z \\u')

# ================= 1 исходные данные =================
section(1, "Исходные данные для технологического проектирования")
k0, k1 = K.K0, B.K1
para("Технологическая карта разработана по заданию № 4 на проектирование технологических процессов земляных "
     "работ, вариант исходных данных 2, и методическим указаниям «Технологическая карта на земляные работы» [1]. "
     "Строительная площадка размером 500 × 300 м разбита на 15 квадратов со стороной a = 100 м. В зоне "
     "квадрата 8 располагается котлован под подземную часть здания. Исходные данные приведены в таблице 1.1, "
     "рабочие отметки вершин квадратов - в таблице 1.2.")
table("Исходные данные (задание № 4, вариант 2)",
      [("Показатель", "Значение")],
      [("Размеры строительной площадки", "500 × 300 м, квадраты 100 × 100 м"),
       ("Грунт", "супесь (по указанию преподавателя)"),
       ("Глубина котлована от планировочной отметки Нк", "2,4 м"),
       ("Высота фундаментной плиты Нф.п", "400 мм"),
       ("Толщина бетонной подготовки hб.п", "150 мм"),
       ("Толщина подсыпки hподс", "100 мм (щебень)"),
       ("Расстояние до карьера, отвала", "10 км"),
       ("Вариант размещения здания", "8 - центр здания в центре квадрата 8"),
       ("Здание по осям (схема здания)", "Г-образное 60 × 30 м с вырезом 42 × 12 м"),
       ("Стены подвала / вылет плиты / выход подготовки / до откоса", "400 / 500 / 200 / 500 мм"),
       ("Толщина растительного слоя hр.сл ([1], с. 11)", "0,2 м"),
       ("Коэффициент первоначального разрыхления kр (прил. 2: 1,12-1,17)", "1,15"),
       ("Коэффициент остаточного разрыхления kо.р (прил. 2: 1,03-1,05)", "1,04")],
      [95, 70], align="ll")
rows = []
for j in range(R.NY + 1):
    rows.append(tuple(f"h{R.vnum(i, j)} = {sgn(R.H[R.vnum(i, j) - 1])}" for i in range(R.NX + 1)))
table("Рабочие отметки вершин квадратов, м (h = Hкр - Hчёрн: «+» - насыпь, «-» - выемка)",
      [], rows, [27.5] * 6, head_rows=0)
para("Вершины пронумерованы построчно слева направо, сверху вниз. Рабочие отметки заданы, поэтому черные и "
     "красные отметки в работе не вычисляются.")

# ================= 2 объёмы земляных работ =================
section(2, "Определение объёмов земляных работ")
h2("2.1 Объёмы планировочных работ")


def zp_rows(P):
    rows = []
    for k, z in enumerate(sorted(P["ZP"], key=lambda z: (z["p1"], z["p2"])), start=1):
        h1_, h2_ = P["H"][z["p1"] - 1], P["H"][z["p2"] - 1]
        side = "горизонт." if z["p2"] == z["p1"] + 1 else "вертик."
        rows.append((k, f"h{z['p1']} - h{z['p2']}", side, sgn(h1_, 4 if P is B.P1 else 2),
                     sgn(h2_, 4 if P is B.P1 else 2), num(z["x1"]), num(z["x2"])))
    return rows


para("Линия нулевых работ (ЛНР) проходит через точки с нулевой рабочей отметкой на сторонах квадратов, "
     "вершины которых имеют отметки разного знака. Положение нулевой точки определяется из подобия "
     "треугольников ([1], рис. 4):")
f_zp = formula("x1 = a · |h1| / (|h1| + |h2|),   x2 = a - x1,")
where("a = 100 м - сторона квадрата", "h1, h2 - рабочие отметки вершин стороны квадрата, м",
      "x1, x2 - расстояния от вершин 1 и 2 до точки нулевых работ, м")
t_zp0 = table("Точки нулевых работ (черновик)",
              [("№", "Сторона", "Положение", "h1, м", "h2, м", "x1, м", "x2, м")], zp_rows(R.P0),
              [10, 28, 28, 25, 25, 25, 24])
para(f"По точкам нулевых работ проведена ЛНР (лист 1 графической части). Она делит площадку на зону "
     f"планировочной выемки (ПВ) и планировочной насыпи (ПН). Квадраты, пересечённые ЛНР, делятся на фигуры: "
     f"выемка - номер квадрата, насыпь - номер со штрихом. Объём грунта в фигуре:")
f_v = formula("V = F · hср,")
where("F - площадь фигуры, м²",
      "hср - среднее абсолютных значений рабочих отметок вершин фигуры, м (точки нулевых работ входят "
      "с нулевой отметкой)")


def vol_rows(P):
    rows, sv, sn, snk = [], 0, 0, 0
    for g in sorted(P["FIGS"], key=lambda g: (g["sq"], g["sign"])):
        cut = g["sign"] < 0
        rows.append((g["name"], num(g["F"]), num(g["hcp"], 4), num(g["V"]) if cut else "",
                     "" if cut else num(g["V"]), "" if cut else num(g["V"] / B.K_OR)))
        if cut: sv += g["V"]
        else: sn += g["V"]; snk += g["V"] / B.K_OR
    rows.append(("Σ", num(sum(g["F"] for g in P["FIGS"])), "", num(sv), num(sn), num(snk)))
    return rows, sv, sn


rows, sv0, sn0 = vol_rows(R.P0)
t_v0 = table("Ведомость объёмов работ по вертикальной планировке, черновик (форма табл. 1 [1])",
             [("№ фигуры", "F, м²", "hср, м", "Vв.геом, м³", "Vн.геом, м³", "Vн / kо.р, м³")], rows,
             [22, 30, 25, 30, 30, 28], bold_last=True)


def slope_text(P, tag):
    s_v, s_n = P["SLOPES"][-1], P["SLOPES"][1]
    para(f"Объёмы откосов, устраиваемых по контуру планируемой площадки ({tag}), - по формуле ([1], с. 10):")
    if tag == "черновик":
        f_ = formula("ΣVв(н) = (Σhв(н) / n)² · Σlв(н) · m / 2,")
        where("Σhв(н) - сумма абсолютных рабочих отметок вершин контура в зоне выемки (насыпи), м",
              "n - количество этих отметок", "Σlв(н) - длина контура площадки в зоне выемки (насыпи), м",
              "m = 0,25 - коэффициент откоса (супесь, высота до 1,5 м, прил. 3 [1])")
    for nm, s_ in (("выемки", s_v), ("насыпи", s_n)):
        para(f"Откосы {nm}: Σh = {num(s_['sum_h'], 4)} м, n = {s_['n']}, Σl = {num(s_['L'])} м; "
             f"ΣV = ({num(s_['sum_h'], 4)} / {s_['n']})² · {num(s_['L'])} · 0,25 / 2 = {num(s_['V'])} м³.")


slope_text(R.P0, "черновик")
para(f"Итого по планировке (черновик): ΣVв.геом = {num(sv0)} + {num(R.P0['SLOPES'][-1]['V'])} = "
     f"{num(R.P0['VV'])} м³; ΣVн.геом = {num(sn0)} + {num(R.P0['SLOPES'][1]['V'])} = {num(R.P0['VN'])} м³; "
     f"ΣVн.геом / kо.р = {num(R.P0['VN'] / B.K_OR)} м³.")

h2("2.2 Контуры, размеры и геометрический объём котлована")
hc = k0["H_CORNERS"]
para(f"ЛНР пересекает котлован (лист 1), поэтому его принадлежность к зоне выемки или насыпи определяется "
     f"по среднему арифметическому рабочих отметок по углам котлована ([1], с. 12). Отметки углов бровки "
     f"(при глубине Нк) найдены линейной интерполяцией в квадрате 8: "
     f"{'; '.join(sgn(v, 3) for v in hc)} м; среднее {sgn(k0['H_MEAN'], 4)} м > 0 - котлован в зоне насыпи. "
     f"Грунт из котлована отрывается от естественной поверхности, и фактическая глубина котлована "
     f"определяется с учётом рабочей отметки в центре котлована hр = {num(k0['HR'], 4)} м (среднее отметок "
     f"вершин квадрата 8):")
formula("hк = Нк - hр.сл - hр,")
where("Нк = 2,4 м - глубина котлована по заданию", "hр.сл = 0,2 м - толщина растительного слоя ([1], с. 11)")
para(f"hк = 2,4 - 0,2 - {num(k0['HR'], 4)} = {num(k0['HK'], 4)} м. Крутизна откосов котлована - по прил. 3 [1]: "
     f"супесь при глубине до 3 м - 1:0,67, m = 0,67; заложение откоса l = hк · m = {num(k0['HK'], 4)} · 0,67 = "
     f"{num(k0['L_OTK'], 4)} м.")
para(f"Размеры котлована откладываются от осей здания: до наружной грани стены - 0,2 м, до края фундаментной "
     f"плиты A = 0,2 + 0,5 = 0,7 м, до края бетонной подготовки - 0,9 м, до низа откоса - 1,4 м (рабочая зона "
     f"С = 0,7 м не менее 0,6 м, [1], с. 12). Со стороны съезда на дне оставляется место для крана: "
     f"c = 1 + 2r + 1 = 1 + 2 · 2,5 + 1 = 7 м от фундаментной плиты (r = 2,5 м - радиус поворотной платформы "
     f"крана), т. е. {num(K.D_PAN, 1)} м от оси. Котлован имеет Г-образную форму (по форме здания). "
     f"Площади по низу и по верху: Fк.н = {num(k0['F_KN'])} м², Fк.в = {num(k0['F_KV'])} м².")
formula("V = hк / 3 · (Fк.н + Fк.в + √(Fк.н · Fк.в)),")
para(f"V = {num(k0['HK'], 4)} / 3 · ({num(k0['F_KN'])} + {num(k0['F_KV'])} + √({num(k0['F_KN'])} · "
     f"{num(k0['F_KV'])})) = {num(k0['V_OSN'])} м³.")
para(f"Съезд в котлован - один, с двусторонним движением автосамосвалов, ширина bп = 6 м ([1], с. 14), уклон "
     f"i = 0,1 (0,10-0,15 при вывозе грунта самосвалами), коэффициент заложения m' = 1 / i = 10; длина съезда в "
     f"плане от низа откоса hк · m' = {num(k0['L_PAN'])} м. Объём съезда сверх откоса котлована:")
formula("Vс = hк² / 6 · (3bп + 2m · hк · (m' - m) / m') · (m' - m),")
para(f"Vс = {num(k0['HK'], 4)}² / 6 · (3 · 6 + 2 · 0,67 · {num(k0['HK'], 4)} · (10 - 0,67) / 10) · (10 - 0,67) = "
     f"{num(k0['V_PAN'])} м³. Объём котлована Vк = {num(k0['V_OSN'])} + {num(k0['V_PAN'])} = {num(k0['V_K'])} м³.")
para("Объём подземной части здания и обратной засыпки ([1], с. 13-14):")
formula("Vп.ч = Vф.п + Vк.с.п = Fф.п · Нф.п + Fк.с.п · hст,")
formula("Vо.з = Vк - Vп.ч - Vподс - Vб.п - Vг.и - Vст,")
where("Fф.п, Fк.с.п - площади фундаментной плиты и по наружному контуру стен, м²",
      "hст = hк - Нф.п - hб.п - hподс - высота стен ниже естественной поверхности, м",
      "Vподс = Fк.н · hподс, Vб.п = Fб.п · hб.п - объёмы подсыпки и бетонной подготовки, м³",
      "Vг.и, Vст - гидроизоляция и защитная стяжка (в задании толщины не даны, приняты 0)")
para(f"Пазухи засыпаются привозным песком (грунт площадки - супесь, не песок), съезд - грунтом котлована. "
     f"Результаты расчёта котлована по черновику и по окончательным отметкам (п. 2.3) - в таблице 2.5.")


def pit_rows():
    keys = [("Средняя рабочая отметка по углам котлована", "", "H_MEAN", 4),
            ("Рабочая отметка в центре котлована", "hр, м", "HR", 4),
            ("Фактическая глубина котлована", "hк, м", "HK", 4),
            ("Заложение откоса", "l, м", "L_OTK", 4),
            ("Площадь котлована по низу", "Fк.н, м²", "F_KN", 2),
            ("Площадь котлована по верху", "Fк.в, м²", "F_KV", 2),
            ("Объём котлована без съезда", "V, м³", "V_OSN", 2),
            ("Длина съезда в плане", "L, м", "L_PAN", 3),
            ("Объём съезда", "Vс, м³", "V_PAN", 2),
            ("Объём котлована", "Vк, м³", "V_K", 2),
            ("Площадь фундаментной плиты", "Fф.п, м²", "F_FP", 2),
            ("Площадь бетонной подготовки", "Fб.п, м²", "F_BP", 2),
            ("Площадь по наружному контуру стен", "Fк.с.п, м²", "F_ST", 2),
            ("Подсыпка (щебень)", "Vподс, м³", "V_PODS", 2),
            ("Бетонная подготовка", "Vб.п, м³", "V_BP", 2),
            ("Фундаментная плита", "Vф.п, м³", "V_FP", 2),
            ("Высота стен ниже естественной поверхности", "hст, м", "H_STEN", 4),
            ("Стены подвала", "Vк.с.п, м³", "V_KSP", 2),
            ("Подземная часть здания", "Vп.ч, м³", "V_PCH", 2),
            ("Засыпка пазух (привозной песок)", "Vп.с, м³", "V_PAZ", 2),
            ("Засыпка съезда (грунт котлована)", "Vс, м³", "V_S", 2),
            ("Обратная засыпка всего", "Vо.з, м³", "V_OZ", 2)]
    return [(t, s, num(k0[key], n), num(k1[key], n)) for t, s, key, n in keys]


TAB_PIT = None   # номер таблицы котлована - в разделе 2.3 (после баланса)

h2("2.3 Сводный баланс грунта")
b0, b1 = B.B0, B.B1


def balance_table(b, k, tag):
    return table(f"Сводная ведомость объёмов разрабатываемого грунта, {tag} (форма табл. 2 [1])",
                 [("Вид работы", "Выемка Vв, м³", "Насыпь Vн, м³", "Vн / kо.р, м³")],
                 [("В призмах (планировка с откосами)", num(b["VV_PL"]), num(b["VN_PL"]), num(b["VN_PL"] / B.K_OR)),
                  ("В котловане (Vк; засыпка съезда Vс)", num(b["V_K"]), num(b["V_S"]), num(b["V_S"] / B.K_OR)),
                  ("Суммарные объёмы", num(b["SUM_V"]), num(b["SUM_N"]), num(b["SUM_NK"])),
                  ("Баланс Vв - Vн / kо.р", num(b["BALANCE"]), "", f"{num(b['PCT'])} %")],
                 [70, 32, 32, 31], align="lccc", bold_last=True)


para("Для соблюдения нулевого баланса земляных масс разность между суммарными объёмами насыпей и выемок не "
     "должна превышать 5 % ([1], с. 14). Пазухи котлована засыпаются привозным песком, поэтому в насыпи "
     "котлована учитывается только засыпка съезда грунтом котлована.")
t_b0 = balance_table(b0, k0, "черновик")
para(f"Баланс отрицательный: {num(b0['BALANCE'])} м³ ({num(b0['PCT'])} % > 5 %) - на площадке преобладает "
     f"насыпь, требуется перерасчёт средней отметки планировки на величину:")
formula("Δh = (ΣVв - ΣVн / kо.р) / Fпл,")
where("Fпл - площадь строительной площадки без площади котлована по верху (котлован в насыпи), м²")
para(f"Fпл = 150 000 - {num(k0['F_KV'])} = {num(b0['F_PL'])} м²; Δh = {num(b0['BALANCE'])} / {num(b0['F_PL'])} = "
     f"{num(b0['DH_EXACT'], 6)} ≈ {num(b0['DH'], 4)} м (понижение планировочной отметки). На Δh скорректированы "
     f"все рабочие отметки: h' = h + Δh (табл. 2.4).")
rows = []
for j in range(R.NY + 1):
    rows.append(tuple(f"h{R.vnum(i, j)} = {sgn(B.H1[R.vnum(i, j) - 1], 4)}" for i in range(R.NX + 1)))
table("Окончательные рабочие отметки h' = h + Δh, м", [], rows, [27.5] * 6, head_rows=0, size=11)
para("По скорректированным отметкам заново определены точки нулевых работ (табл. 2.6), объёмы планировочных "
     "работ (табл. 2.7), откосы, котлован (табл. 2.5) и сводный баланс (табл. 2.8).")
# таблица котлована - здесь, после окончательных отметок
t_pit = table("Параметры и объёмы котлована",
              [("Показатель", "Обозн.", "Черновик", "Окончат.")], pit_rows(), [80, 27, 29, 29], align="lccc")
table("Точки нулевых работ по окончательным отметкам",
      [("№", "Сторона", "Положение", "h1, м", "h2, м", "x1, м", "x2, м")], zp_rows(B.P1),
      [10, 28, 28, 25, 25, 25, 24])
rows, sv1, sn1 = vol_rows(B.P1)
table("Ведомость объёмов работ по вертикальной планировке по окончательным отметкам (форма табл. 1 [1])",
      [("№ фигуры", "F, м²", "hср, м", "Vв.геом, м³", "Vн.геом, м³", "Vн / kо.р, м³")], rows,
      [22, 30, 25, 30, 30, 28], bold_last=True)
slope_text(B.P1, "окончательные отметки")
para(f"Итого по планировке: ΣVв.геом = {num(B.P1['VV'])} м³; ΣVн.геом = {num(B.P1['VN'])} м³. Котлован "
     f"остаётся в зоне насыпи, глубина пересчитана: hк = {num(k1['HK'], 4)} м, Vк = {num(k1['V_K'])} м³ (табл. 2.5).")
balance_table(b1, k1, "по окончательным отметкам")
para(f"Баланс {num(b1['BALANCE'])} м³ ({num(b1['PCT'])} % ≤ 5 %) - условие соблюдено. Недостающий грунт "
     f"подвозится из карьера ([1], с. 15).")

h2("2.4 Распределение грунта в котловане")
para("Грунт, извлечённый из котлована, распределяется по строительной площадке для укладки в насыпь, "
     "складируется в отвал для обратной засыпки и вывозится при положительном балансе ([1], с. 16). Пазухи "
     "засыпаются привозным песком, поэтому в отвал у котлована складируется только грунт для засыпки съезда; "
     "остальной грунт котлована - в транспортные средства и в насыпь планировки. Баланс земляных работ по "
     "котловану составлен в форме, приведённой на практическом занятии, с содержанием табл. 3 [1] (табл. 2.9).")
rows = [(1, RS.T3_V[0][0], num(RS.T3_V[0][1]), "Отвал (засыпка съезда)", num(RS.T3_OTVAL),
         1, RS.T3_N[0][0], num(RS.T3_N[0][1]), num(RS.T3_N[0][2])),
        (2, RS.T3_V[1][0], num(RS.T3_V[1][1]), "Транспортные средства", num(RS.T3_TRANSPORT),
         2, RS.T3_N[1][0], num(RS.T3_N[1][1]), num(RS.T3_N[1][2])),
        ("", "ΣVв", num(RS.T3_SUM_V), "Σ (с kр)", num(RS.T3_OTVAL + RS.T3_TRANSPORT),
         "", "ΣVн местного грунта", "", num(RS.T3_SUM_N))]
table("Баланс земляных работ по котловану (табл. 3)",
      [("№", "Место разработки грунта", "Объём, м³ (плотн.)", "Место укладки", "Объём с kр, м³",
        "№", "Место устройства насыпи", "Геом. объём, м³", "Потребно V / kо.р, м³")],
      rows, [7, 20, 18, 24, 17, 7, 32, 18, 22], size=10, align="clcccclcc", bold_last=True)
para(f"В отвал ΣVн · kр = {num(RS.T3_SUM_N)} · 1,15 = {num(RS.T3_OTVAL)} м³; в транспортные средства "
     f"(ΣVв - ΣVн) · kр = ({num(RS.T3_SUM_V)} - {num(RS.T3_SUM_N)}) · 1,15 = {num(RS.T3_TRANSPORT)} м³, из них в "
     f"насыпь планировки {num(RS.V_PIT_FILL)} м³ в плотном теле, вывоз {num(RS.V_LISH)} м³ (баланс отрицательный). "
     f"Из карьера привозятся: песок для засыпки пазух {num(RS.T3_PRIVOZ[0][1])} м³ и щебень для подсыпки "
     f"{num(RS.T3_PRIVOZ[1][1])} м³ (в плотном теле, V / kо.р).")

h2("2.5 Распределение земляных масс на площадке. Картограмма")
para("Распределение земляных масс выполнено с учётом наименьших моментов перемещений ([1], с. 19): минимум "
     "суммы произведений объёмов на расстояния между центрами тяжести фигур выемки и насыпи. Сначала грунт "
     "выемок перемещается в насыпь землеройно-транспортными машинами, затем недостающий грунт в насыпи "
     "восполняется грунтом из котлована (автосамосвалы, от центра котлована); остаток - недостача, подвоз из "
     "карьера. Объёмы фигур выемки - в плотном теле, насыпи - с учётом kо.р. Машина назначается по "
     "дальности ([1], с. 22): до 70 м - бульдозер, дальше - скрепер. Перемещения нанесены на картограмму "
     "(лист 4), стрелки землеройно-транспортных машин не пересекаются.")
rows = [(a, b, num(v), num(l, 1), m) for a, b, v, l, m in RS.MOVES]
rows.append(("Недостача", ", ".join(RS.FILL[j]["name"] for j in RS.SHORT), num(sum(RS.SHORT.values())), "",
             "из карьера"))
table("Перемещения грунта (баланс распределения земляных масс, табл. 4 [1])",
      [("Откуда (фигура)", "Куда (фигура)", "V, м³", "l, м", "Машина")], rows, [35, 30, 35, 30, 35])

h2("2.6 Определение средней дальности перемещения грунта")
para("Средняя дальность перемещения грунта - расстояние между центрами тяжести равновеликих по объёму "
     "участков выемки и насыпи - определена методом статических моментов отдельно для землеройно-"
     "транспортных машин и для автосамосвалов ([1], с. 19-21):")
formula("Mx = x · V,   My = y · V,")
formula("xц = ΣMx / ΣV,   yц = ΣMy / ΣV,")
formula("Lср = √((xц.в - xц.н)² + (yц.в - yц.н)²).")


def mean_table(data, title, labels):
    rs, rd, cs, cd, L = data
    rows = [(labels[0], nm, num(v), num(x), num(y), num(mx, 1), num(my, 1)) for nm, v, x, y, mx, my in rs] + \
           [(labels[1], nm, num(v), num(x), num(y), num(mx, 1), num(my, 1)) for nm, v, x, y, mx, my in rd]
    table(title, [("Зона", "Фигура", "V, м³", "x, м", "y, м", "Mx = x·V", "My = y·V")], rows,
          [22, 22, 26, 20, 20, 28, 27], size=11)
    para(f"Центр тяжести «{labels[0].lower()}»: ({num(cs[0])}; {num(cs[1])}) м; «{labels[1].lower()}»: "
         f"({num(cd[0])}; {num(cd[1])}) м. Lср = √(({num(cs[0])} - {num(cd[0])})² + ({num(cs[1])} - "
         f"{num(cd[1])})²) = {num(L)} м.")


mean_table(RS.T5, "Средняя дальность перемещения грунта скрепером (бульдозером): выемка - насыпь (табл. 5 [1])",
           ("Выемка", "Насыпь"))
mean_table(RS.T6, "Средняя дальность перемещения грунта автосамосвалом: котлован - насыпь (табл. 6 [1])",
           ("Котлован", "Насыпь"))
para(f"Проверка аналитическим методом Lср = Σ(V · l) / ΣV ([1], с. 19): для землеройно-транспортных машин "
     f"{num(RS.L_ANALYT['ЗТМ'])} м, для автосамосвалов {num(RS.L_ANALYT['самосвал'])} м.")

# ================= 3 комплект машин =================
section(3, "Назначение комплекта машин")
para(f"Наименование и краткая характеристика грунта - {MS.GRUNT.lower()} ([2], техническая часть, табл. 1, п. 22); "
     f"группа грунта в зависимости от трудности разработки: одноковшовыми экскаваторами - {MS.GR_EXC}, "
     f"скреперами - {MS.GR_SCR}, бульдозерами - {MS.GR_BUL}; средняя плотность в естественном залегании - "
     f"{MS.RHO} кг/м³.")
h2("3.1 Машины для вертикальной планировки строительной площадки")
para(f"Средняя дальность перемещения грунта Lср = {num(MS.L_SR)} м (табл. 2.11). При дальности от 150 до 1000 м "
     f"применяют прицепные скреперы с ковшом до 10 м³ ([1], с. 22; прил. 5: ковш 6 м³ - 100-350 м). Для "
     f"технико-экономического сравнения рассмотрены два варианта (табл. 7 [1]): ведущая машина - бульдозер "
     f"(вариант 1) и скрепер (вариант 2). Супесь разрабатывается скреперами без предварительного рыхления "
     f"(прил. 5), трактор-рыхлитель включён в оба варианта по форме табл. 7 [1]. Нормы времени - по ЕНиР [2] "
     f"(табл. 3.1); добавка на каждые следующие 10 м сверх расстояния, указанного в норме, принята "
     f"пропорционально. Работа - в {MS.M_SM} смены по 8 ч, срок планировки t = {MS.T_DN} дней (5-20 дней, "
     f"одинаковый для вариантов).")
norm_rows = []
for m in (MS.BUL_VED, MS.SKR_VED, MS.RYKHL, MS.BUL_RAZ, MS.KATOK, MS.TOLKACH):
    role, brand, src, a, b, L0, note = m
    hv = "Нвр скрепера / 3" if a is None else (num(a) + (f" + {num(b)} на 10 м сверх {L0} м" if b else ""))
    norm_rows.append((role, brand, src, hv))
table("Нормы времени машин, маш.-ч на 100 м³ (ЕНиР Е2-1 [2])",
      [("Машина", "Марка", "Обоснование", "Нвр")], norm_rows, [32, 48, 45, 40], size=11, align="llll")
para("Дневная выработка, количество машино-смен и количество машин ([1], с. 25-27):")
formula("Vдн = 100 · N · 8 / Нвр,")
formula("Mсм = Нвр · V / (100 · 8),")
formula("nвед = Mсм.вед / (m · t),")
formula("n = nвед · Vдн.вед / Vдн.об,")
where("N = m = 2 - число смен в сутки", "V - объём работ машины, м³ (ведущая машина и толкач - с kр, рыхлитель - "
      "в плотном теле, бульдозер на разравнивании и каток - объём насыпи)")


def var_rows(v, title):
    rows = [title]
    for r in v["rows"]:
        n_ = str(r["n"]) if "n_calc" not in r else f"{num(r['n_calc'], 3)} → {r['n']}"
        rows.append((r["mach"][0], r["mach"][1].split(",")[0], num(r["V"]), num(r["H"], 4), num(r["M"]),
                     num(r["Vdn"]), n_))
    rows.append(f"Итого: Mсм.вед + ΣMсм.об = {num(v['SUM_M'])} маш.-см; на 1 м³ - {num(v['M_PER_V'], 5)} маш.-см; "
                f"V / ΣMсм = {num(v['V_PER_M'])} м³/маш.-см")
    return rows


rows = var_rows(MS.VAR1, "Вариант 1 - ведущая машина бульдозер") + var_rows(MS.VAR2, "Вариант 2 - ведущая машина скрепер")
table("Варианты с различными ведущими машинами (табл. 7 [1])",
      [("Машина", "Марка", "V, м³", "Нвр, маш.-ч", "Mсм, маш.-см", "Vдн, м³/дн", "n, шт.")], rows,
      [33, 32, 22, 18, 20, 20, 20], size=10, align="llccccc")
para(f"Затраты машинного времени на 1 м³: вариант 1 - {num(MS.VAR1['M_PER_V'], 5)} маш.-см, вариант 2 - "
     f"{num(MS.VAR2['M_PER_V'], 5)} маш.-см. Принимается вариант {MS.VAR_OK} - скреперный комплект: "
     f"{MS.VAR2['rows'][0]['n']} прицепных скрепера ДЗ-20 с трактором-толкачом Т-180, бульдозер-рыхлитель "
     f"ДП-22С, бульдозер ДЗ-35С, каток ДУ-16В; продолжительность работы ведущих машин "
     f"{num(MS.VAR2['DAYS'], 1)} дня.")

h2("3.2 Машины для разработки грунта в котловане")
para(f"Объём котлована Vк = {num(MS.V_K)} м³; по прил. 6 [1] рекомендуемая ёмкость ковша экскаватора - "
     f"{' или '.join(MS.KOVSH_PRIL6)} м³. Котлован глубиной до 3 м разрабатывается одноковшовым экскаватором, "
     f"оборудованным обратной лопатой, с погрузкой в автосамосвалы ([2], § Е2-1-11). Сравнены три экскаватора "
     f"(табл. 3.3); принят экскаватор с наибольшей производительностью.")
rows = [(e[0], e[1], num(e[2]), num(e[3], 1), num(MS.P_SM[e[0]]), num(e[5], 1), num(e[6], 1), num(e[7], 1), e[8])
        for e in MS.EXC]
table("Экскаваторы с обратной лопатой (ЕНиР Е2-1-11 [2], грунт I группы)",
      [("Марка", "Привод", "Ковш, м³", "Нвр, маш.-ч/100 м³", "Псм, м³/смену", "Hк max, м", "Rmax, м",
        "Hвыгр, м", "Недобор, см")], rows, [19, 29, 14, 19, 19, 15, 15, 16, 19], size=10)
para(f"Принят экскаватор {MS.EX_NAME} (гидравлический, ковш {num(MS.V_KOV)} м³). Допускаемый недобор грунта "
     f"в основании - {MS.NEDOBOR} см (прил. 7 [1]); глубина копания экскаватором {num(MS.HK, 4)} - "
     f"{num(MS.NEDOBOR / 100)} = {num(MS.H_KOP, 4)} м < Hк max = {num(MS.EX_HK, 1)} м. "
     f"Затраты машинного времени Mсм = {num(MS.H_EX, 1)} · {num(MS.V_K)} / 800 = {num(MS.M_EX)} маш.-см, "
     f"продолжительность разработки одним экскаватором в 2 смены - {num(MS.DAYS_EX, 1)} дня.")
para("Количество автосамосвалов для обслуживания экскаватора ([1], с. 28-29):")
formula("Vгр.к = Vковша · kнап / kр,   Mковша = Vгр.к · γ,")
formula("n = min(Па / Mковша; Va / Vковша),   6 ≤ n ≤ 11,")
formula("Tцикла = Ту.п + Tп + Tпр.гр + Tпр.пор + Ту.р + Tр + Tм,")
formula("N ≥ Tцикла / Tп,   Нвр.а = Tцикла / (Va.р · 60),")
where(f"kнап = {num(MS.K_NAP, 1)} - коэффициент наполнения ковша обратной лопаты (табл. 8 [1]: 0,8-1,0)",
      f"γ = {num(MS.RHO / 1000, 2)} т/м³ - плотность грунта", "Па, Va - грузоподъёмность и объём кузова",
      f"Tп = Va.пл · Нвр / 100 · 60 - время погрузки, Va.пл = n · Vгр.к",
      f"Tпр.гр = L / Uгр · 60, Tпр.пор = L / Uпор · 60 - пробег, L = {num(MS.L_TR, 0)} км (задание: расстояние "
      f"до карьера, отвала)", "Ту.п = 0,3 мин, Ту.р = 0,6 мин, Tр = 1 мин, Tм = 2 мин")
para(f"Vгр.к = {num(MS.V_KOV)} · {num(MS.K_NAP, 1)} / 1,15 = {num(MS.V_GRK, 4)} м³; Mковша = {num(MS.V_GRK, 4)} · "
     f"{num(MS.RHO / 1000, 2)} = {num(MS.M_KOV, 4)} т. Расчёт по автосамосвалам прил. 8 [1] - в табл. 3.4.")
rows = [(t["name"], num(t["Va"], 1), num(t["Pa"], 1), t["n"], "да" if t["ok"] else "нет", num(t["Tc"]),
         t["N"], num(t["Ha"], 4)) for t in MS.TRUCKS]
table("Подбор автосамосвалов (прил. 8 [1])",
      [("Модель", "Va, м³", "Па, т", "n", "6 ≤ n ≤ 11", "Tцикла, мин", "N, шт.", "Нвр.а, маш.-ч/м³")], rows,
      [40, 16, 16, 12, 20, 22, 16, 23], size=10, align="lccccccc")
tr = MS.TR
para(f"Из автосамосвалов, удовлетворяющих условию 6 ≤ n ≤ 11, принят {tr['name']} с наименьшей нормой времени на "
     f"1 м³: n = {tr['n']} ковшей, Tп = {num(tr['Tp'])} мин, Tцикла = {num(tr['Tc'])} мин, N = {num(tr['N_calc'])} "
     f"→ {tr['N']} автосамосвалов.")

# ================= графическая часть и литература =================
p = para("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", bold=True, align=C, indent=False, space_after=12)
p.paragraph_format.page_break_before = True
p.style = doc.styles["Heading 1"]
p.paragraph_format.page_break_before = True
p.alignment = C
p.paragraph_format.first_line_indent = Cm(0)
for k, s_ in enumerate((
        "Технологическая карта на земляные работы : методические указания к практическим занятиям для "
        "обучающихся по направлению подготовки 08.05.01 Строительство / сост. А.А. Лапидус, А.С. Комиссарова, "
        "Е.М. Пугач ; НИУ МГСУ, каф. технологии и организации строительного производства. - Москва : "
        "Издательство МИСИ - МГСУ, 2018. - 48 с.",
        "ЕНиР. Сборник Е2. Земляные работы. Выпуск 1. Механизированные и ручные земляные работы. - Москва : "
        "Стройиздат, 1988.",
        "СП 45.13330.2012. Земляные сооружения, основания и фундаменты. Актуализированная редакция "
        "СНиП 3.02.01-87.",
        "СНиП 12-03-2001. Безопасность труда в строительстве. Часть 1. Общие требования.",
        "СНиП 12-04-2002. Безопасность труда в строительстве. Часть 2. Строительное производство.",
        "ГОСТ Р 21.101-2020. Система проектной документации для строительства. Основные требования к "
        "проектной и рабочей документации."), start=1):
    para(f"{k}. {s_}")

p = para("ПРИЛОЖЕНИЕ А", bold=True, align=C, indent=False)
p.style = doc.styles["Heading 1"]
p.alignment, p.paragraph_format.first_line_indent = C, Cm(0)
para("(обязательное)", align=C, indent=False)
para("Графическая часть", bold=True, align=C, indent=False, space_after=12)
para("Графическая часть выполнена на листах формата А3 ([1], разд. 3) и приведена на следующих страницах:")
for title_, k in SHEETS:
    para(f"{title_}.", indent=True)
IMG.mkdir(parents=True, exist_ok=True)
for title_, k in SHEETS:
    pdf = HERE / "Чертежи" / f"{title_}.pdf"
    png = IMG / f"list{k}.png"
    pymupdf.open(pdf)[0].get_pixmap(dpi=200).save(png)
    s_ = doc.add_section(WD_SECTION.NEW_PAGE)
    s_.orientation = WD_ORIENT.LANDSCAPE
    s_.page_width, s_.page_height = Mm(420), Mm(297)
    s_.left_margin = s_.right_margin = s_.top_margin = s_.bottom_margin = Mm(0)
    s_.header_distance = s_.footer_distance = Mm(0)
    s_.different_first_page_header_footer = False
    s_.footer.is_linked_to_previous = False
    s_.header.is_linked_to_previous = False
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent, p.paragraph_format.line_spacing = Cm(0), 1.0
    p.alignment = C
    p.add_run().add_picture(str(png), width=Mm(419), height=Mm(296))

doc.save(OUT)

# Word: обновить содержание и поля, сохранить, выгрузить PDF
pdf_out = OUT.with_suffix(".pdf")
ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; $w.DisplayAlerts = 0; "
      f"try {{ $d = $w.Documents.Open('{OUT}'); $d.TablesOfContents.Item(1).Update(); $d.Fields.Update() | Out-Null; "
      f"$d.TablesOfContents.Item(1).Update(); $d.Save(); $d.ExportAsFixedFormat('{pdf_out}', 17); "
      f"$n = $d.ComputeStatistics(2); $d.Close(); Write-Output $n }} finally {{ $w.Quit() }}")
r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True)
print(OUT.name, "- страниц:", r.stdout.strip() or "?", r.stderr.strip()[:300])
